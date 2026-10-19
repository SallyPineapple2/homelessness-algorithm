"""
Graphs of who would get the housing: VI-SPDAT vs. Gemini, for every characteristic.

Figure 1  Race, gender, city, and disclosure: each group's share of housing referrals (score 8+)
          and of scarce housing (top 25% of records), against the equal share a fair tool gives.
Figure 2  The 16 VI-SPDAT traits: among people NOT truly high-need, the share Gemini still refers,
          with vs. without the trait (the VI-SPDAT refers none of them).
Figure 3  The 16 traits: among truly high-need people, the share each tool refers, with vs.
          without the trait.

Reads the results of scripts/allocation_shares.py and scripts/allocation_by_trait.py.
APA 7 style: grayscale, Times New Roman, groups identified by axis labels and marker shape.

Writes: exports/figures/allocation_fig1_demographics.png, allocation_fig2_over_referral.png,
        allocation_fig3_caught.png, and exports/allocation_figures.docx
Run:    python scripts/figures_allocation.py
"""

from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
RES = ROOT / "data" / "ai" / "results"
FIG = ROOT / "exports" / "figures"
FIG.mkdir(parents=True, exist_ok=True)
AI = "Gemini"

shares = pd.read_csv(RES / "allocation_shares_gemini.csv")
traits = pd.read_csv(RES / "allocation_by_trait_gemini.csv")

plt.rcParams.update({"font.family": "Times New Roman", "font.size": 10, "axes.spines.top": False,
                     "axes.spines.right": False, "savefig.dpi": 300, "savefig.bbox": "tight"})


def fp(p):
    return "< .001" if p < 0.001 else f"= {p:.3f}".replace("0.", ".")


# ---------------- Figure 1: demographics vs. equal share ----------------
factors = ["Gender", "Race", "City", "Disclosure"]
heights = [len(shares[(shares.factor == f) & (shares.group != "p (permutation)")]) for f in factors]
fig, axes = plt.subplots(len(factors), 1, figsize=(7.0, 1.0 + 0.42 * sum(heights) + 0.6 * len(factors)),
                         gridspec_kw={"height_ratios": heights})
MARKS = [("VI-SPDAT PSH share", "VI-SPDAT, referred", "s", "black"),
         ("VI-SPDAT top share", "VI-SPDAT, scarce", "s", "white"),
         (f"{AI} PSH share", f"{AI}, referred", "o", "black"),
         (f"{AI} top share", f"{AI}, scarce", "o", "white")]
offsets = [-0.27, -0.09, 0.09, 0.27]
for ax, f in zip(axes, factors):
    sub = shares[(shares.factor == f) & (shares.group != "p (permutation)")].reset_index(drop=True)
    pr = shares[(shares.factor == f) & (shares.group == "p (permutation)")].iloc[0]
    y = np.arange(len(sub))[::-1]
    fair = 100 * sub.equal_share.iloc[0]
    ax.axvline(fair, color="black", linestyle="--", linewidth=1, zorder=1)
    for (col, name, mk, face), off in zip(MARKS, offsets):
        ax.scatter(100 * sub[col], y + off, marker=mk, s=34, facecolors=face, edgecolors="black",
                   linewidths=1, zorder=3, label=name)
    ax.set_yticks(y)
    ax.set_yticklabels(sub.group)
    ax.tick_params(axis="y", length=0)
    lo = min(100 * sub[[c for c, *_ in MARKS]].min().min(), fair) - 3
    hi = max(100 * sub[[c for c, *_ in MARKS]].max().max(), fair) + 3
    ax.set_xlim(lo, hi)
    ax.set_title(f"{f} (fair share {fair:.0f}%) — {AI}: referred p {fp(pr[f'{AI} PSH share'])}, "
                 f"scarce p {fp(pr[f'{AI} top share'])}", loc="left", fontsize=10)
axes[-1].set_xlabel("Share of housing going to each group (%)")
handles, labels = axes[0].get_legend_handles_labels()
fig.legend(handles, labels, loc="lower center", ncol=4, frameon=False, bbox_to_anchor=(0.5, -0.04))
fig.tight_layout(h_pad=1.4)
fig.savefig(FIG / "allocation_fig1_demographics.png")
plt.close(fig)


# ---------------- dumbbell helper for Figures 2 and 3 ----------------
def dumbbell(ax, data, with_col, without_col, sig_col, title):
    d = data.reset_index(drop=True)
    y = np.arange(len(d))[::-1]
    for yi, (_, r) in zip(y, d.iterrows()):
        a, b = r[with_col], r[without_col]
        if pd.notna(a) and pd.notna(b):
            ax.plot([100 * b, 100 * a], [yi, yi], color="0.55", linewidth=1.4, zorder=1)
        if pd.notna(b):
            ax.scatter(100 * b, yi, marker="o", s=34, facecolors="white", edgecolors="black", linewidths=1, zorder=3)
        if pd.notna(a):
            ax.scatter(100 * a, yi, marker="o", s=34, facecolors="black", edgecolors="black", zorder=3)
        if sig_col and pd.notna(r[sig_col]) and r[sig_col] < 0.05:
            ax.text(103, yi, "*", va="center", fontsize=12)
    ax.set_yticks(y)
    ax.set_yticklabels(d.trait)
    ax.tick_params(axis="y", length=0)
    ax.set_xlim(-3, 108)
    ax.set_title(title, loc="left", fontsize=10)


short = {"Mental Health / Tri-Morbidity": "Mental Health / Tri-Morbidity", "Duration and episodes": "Long homelessness",
         "Housing status": "Unsheltered or unstable housing", "Money Management": "Money management (finances)",
         "Physical Health": "Physical health (illness)"}
traits["trait"] = traits["trait"].map(lambda t: short.get(t, t))

# ---------------- Figure 2: over-referral ----------------
t2 = traits.sort_values("ai_over_diff", ascending=True, na_position="first")
fig, ax = plt.subplots(figsize=(7.0, 5.6))
dumbbell(ax, t2.iloc[::-1], "ai_over_with", "ai_over_without", "p_over_holm",
         f"{AI}: referred to housing although NOT truly high-need")
ax.set_xlabel("Share referred (%)")
ax.scatter([], [], marker="o", s=34, facecolors="black", edgecolors="black", label="Has the characteristic")
ax.scatter([], [], marker="o", s=34, facecolors="white", edgecolors="black", label="Does not have it")
ax.legend(loc="lower center", ncol=2, frameon=False, bbox_to_anchor=(0.5, -0.2))
fig.tight_layout()
fig.savefig(FIG / "allocation_fig2_over_referral.png")
plt.close(fig)

# ---------------- Figure 3: caught ----------------
t3 = traits.sort_values("ai_caught_diff", ascending=False, na_position="last")
fig, axes = plt.subplots(1, 2, figsize=(7.5, 5.6), sharey=True)
dumbbell(axes[0], t3, "vis_caught_with", "vis_caught_without", None, "VI-SPDAT")
dumbbell(axes[1], t3, "ai_caught_with", "ai_caught_without", "p_caught_holm", AI)
for ax, letter in zip(axes, "AB"):
    ax.set_xlabel("Truly high-need people referred (%)")
    ax.text(-0.02, 1.02, letter, transform=ax.transAxes, fontsize=11, fontweight="bold", ha="right", va="bottom")
axes[0].scatter([], [], marker="o", s=34, facecolors="black", edgecolors="black", label="Has the characteristic")
axes[0].scatter([], [], marker="o", s=34, facecolors="white", edgecolors="black", label="Does not have it")
fig.legend(*axes[0].get_legend_handles_labels(), loc="lower center", ncol=2, frameon=False, bbox_to_anchor=(0.5, -0.05))
fig.tight_layout()
fig.savefig(FIG / "allocation_fig3_caught.png")
plt.close(fig)

# ---------------- Word file ----------------
doc = Document()
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(doc.sections[0], side, Inches(1))
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(12)


def figure(n, title, path, note):
    p = doc.add_paragraph()
    p.add_run(f"Figure {n}").bold = True
    p.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph()
    p.add_run(title).italic = True
    p.paragraph_format.space_after = Pt(6)
    doc.add_picture(str(path), width=Inches(6.5))
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.add_run("Note. ").italic = True
    p.add_run(note).font.size = Pt(10)
    doc.add_page_break()


figure(1, "Share of Housing Going to Each Demographic Group, VI-SPDAT and Gemini", FIG / "allocation_fig1_demographics.png",
       "Referred: records scoring 8 or more (Permanent Supportive Housing). Scarce: only the top 25% of the 640 "
       "records can be housed; ties at the cutoff are split evenly. Dashed line: the equal share a tool with no "
       "demographic tendency gives, since every group appears equally often. p: permutation test of Gemini's "
       "departure from equal shares (labels shuffled within each base profile, 10,000 shuffles). The VI-SPDAT "
       "does not use race, gender, or city; its disclosure split differs from equal at p < .001.")
figure(2, "People Referred to Housing Although Not Truly High-Need, With and Without Each Characteristic",
       FIG / "allocation_fig2_over_referral.png",
       "Gemini only; the VI-SPDAT referred none of these records. Characteristics are VI-SPDAT indicators written "
       "into each base profile. Truly high-need: true score 8 or more. * significant after Holm correction "
       "(permutation test reassigning the characteristic across the 32 base profiles). Mental Health occurs only "
       "on truly high-need profiles, so it has no value here.")
figure(3, "Truly High-Need People Referred to Housing, With and Without Each Characteristic",
       FIG / "allocation_fig3_caught.png",
       "(A) VI-SPDAT. (B) Gemini. * Gemini's difference is significant after Holm correction (permutation test "
       "across the 32 base profiles). Some groups are small: for several characteristics only one truly "
       "high-need base profile (20 records) lacks it.")
doc.save(ROOT / "exports" / "allocation_figures.docx")
print("wrote exports/allocation_figures.docx and three figures")
