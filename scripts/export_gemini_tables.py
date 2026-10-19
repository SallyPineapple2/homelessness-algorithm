"""
Gemini-only bias tables for the paper.

Table A  Gemini's demographic bias tested against itself (the dashboard's original tests):
         paired differences between versions of the same record, Holm-adjusted across the ten
         demographic tests, with the VI-SPDAT's result on the same test for reference.
Table B  Gemini compared with the VI-SPDAT: the gap (Gemini score minus VI-SPDAT score on the same
         record) compared between groups, paired across the 32 base profiles. Holm-adjusted
         within Gemini's own tests (race vs. White, 4 tests).

Writes: exports/gemini_tables.docx
Run:    python scripts/export_gemini_tables.py
"""

import csv
import json
from collections import defaultdict
from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from scipy import stats

ROOT = Path(__file__).resolve().parents[1]
results = json.loads((ROOT / "data" / "ai_results.json").read_text(encoding="utf-8"))
gem = next(m for m in results["models"] if m["key"] == "gemini")
vis = results["vispdat"]

RACES = ["American Indian, Alaska Native, or Indigenous", "Asian or Asian American",
         "Black, African American, or African", "Hispanic/Latino/e/a", "White"]
SHORT = {"American Indian, Alaska Native, or Indigenous": "American Indian or Alaska Native",
         "Asian or Asian American": "Asian", "Black, African American, or African": "Black",
         "Hispanic/Latino/e/a": "Hispanic or Latino"}


def fp(p):
    return "< .001" if p < 0.001 else f"{p:.3f}".lstrip("0")


def short(comp):
    for k, v in SHORT.items():
        comp = comp.replace(k, v)
    return (comp.replace(", California", "").replace(", Ohio", "").replace(", Illinois", "")
                .replace(", Georgia", "").replace("New York City, New York", "New York City")
                .replace("other locations", "other four cities").replace(" vs. Male", " vs. male"))


def holm(ps):
    order = np.argsort(ps)
    adj, run = [0.0] * len(ps), 0.0
    for rank, i in enumerate(order):
        run = max(run, min(1.0, (len(ps) - rank) * ps[i]))
        adj[i] = run
    return adj


# ---------- Table A rows: Gemini against itself ----------
vis_by = {t["comparison"]: t for t in vis["tests"]}
rows_a = []
for t in gem["tests"]:
    v = vis_by[t["comparison"]]
    rows_a.append([
        short(t["comparison"]), str(t["n"]),
        f'{t["mean_diff"]:+.2f} [{t["ci_low"]:+.2f}, {t["ci_high"]:+.2f}]',
        f'{t["t"]:.2f}', f'{t["d_z"]:.2f}', fp(t["p"]), fp(t["p_holm"]),
        "Yes" if t["significant_holm"] else "No",
        f'{v["mean_diff"]:+.2f}' + (" *" if v["significant_holm"] else ""),
    ])

# ---------- Table B rows: Gemini compared with the VI-SPDAT ----------
recs = list(csv.DictReader((ROOT / "data" / "ai" / "results" / "gemini.csv").open(encoding="utf-8-sig")))
gap = defaultdict(list)
for r in recs:
    g = int(r["score"]) - int(r["vispdat_score"])
    gap[("race", r["race"], r["profile"])].append(g)
    gap[("gender", r["gender"], r["profile"])].append(g)
profiles = sorted({r["profile"] for r in recs})


def pp(field, level):
    return np.array([np.mean(gap[(field, level, p)]) for p in profiles])


def paired(a, b):
    d = a - b
    t = stats.ttest_1samp(d, 0)
    ci = stats.t.interval(0.95, len(d) - 1, loc=d.mean(), scale=stats.sem(d))
    return d.mean(), ci, t.statistic, t.pvalue, d.mean() / d.std(ddof=1)


X = np.column_stack([pp("race", r) for r in RACES])
n, k = X.shape
grand = X.mean()
ss_c = n * ((X.mean(0) - grand) ** 2).sum()
ss_s = k * ((X.mean(1) - grand) ** 2).sum()
ss_e = ((X - grand) ** 2).sum() - ss_c - ss_s
F = (ss_c / (k - 1)) / (ss_e / ((k - 1) * (n - 1)))
C = np.eye(k) - 1 / k
S = C @ np.cov(X, rowvar=False) @ C
eps = np.trace(S) ** 2 / ((k - 1) * np.sum(S * S))
p_omni = stats.f.sf(F, (k - 1) * eps, (k - 1) * (n - 1) * eps)
eta = ss_c / (ss_c + ss_e)

pairs = [(SHORT[r], paired(pp("race", r), pp("race", "White"))) for r in RACES[:-1]]
padj = holm([x[1][3] for x in pairs])
gd = paired(pp("gender", "Female"), pp("gender", "Male"))

means = {r: pp("race", r).mean() for r in RACES}
means.update({g: pp("gender", g).mean() for g in ("Male", "Female")})

rows_b = [["Race (all five groups)", "", "", f"F(4, 124) = {F:.2f}", f"η²p = {eta:.3f}", fp(p_omni), "—",
           "Yes" if p_omni < 0.05 else "No"]]
for (name, (d, ci, t, p, dz)), pa in zip(pairs, padj):
    rows_b.append([f"{name} vs. White", f"{means[[r for r in RACES if SHORT.get(r) == name][0]]:.2f} vs. {means['White']:.2f}",
                   f"{d:+.2f} [{ci[0]:+.2f}, {ci[1]:+.2f}]", f"t(31) = {t:.2f}", f"d_z = {dz:.2f}", fp(p), fp(pa),
                   "Yes" if pa < 0.05 else "No"])
d, ci, t, p, dz = gd
rows_b.append(["Female vs. male", f"{means['Female']:.2f} vs. {means['Male']:.2f}", f"{d:+.2f} [{ci[0]:+.2f}, {ci[1]:+.2f}]",
               f"t(31) = {t:.2f}", f"d_z = {dz:.2f}", fp(p), "—", "Yes" if p < 0.05 else "No"])

# ---------- Word file ----------
doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Inches(1))
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(12)


def rule(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:color"), "000000")
    b.append(el)


def apa_table(number, title, heads, rows, widths, note):
    p = doc.add_paragraph()
    p.add_run(f"Table {number}").bold = True
    p.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph()
    p.add_run(title).italic = True
    p.paragraph_format.space_after = Pt(6)
    t = doc.add_table(rows=1, cols=len(heads))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(heads):
        c = t.rows[0].cells[i]
        c.text = ""
        c.paragraphs[0].add_run(h).font.size = Pt(10)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
        rule(c, "top", 8)
        rule(c, "bottom", 4)
    for n_, r in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(v)
            run.font.size = Pt(10)
            if heads[i] == "Significant" and v == "Yes":
                run.bold = True
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            if n_ == len(rows) - 1:
                rule(cells[i], "bottom", 8)
    for row in t.rows:
        for c, w in zip(row.cells, widths):
            c.width = Inches(w)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.add_run("Note. ").italic = True
    p.add_run(note).font.size = Pt(10)
    doc.add_paragraph()


apa_table(
    "A", "Demographic Differences Within Gemini's Own Scores",
    ["Comparison", "n", "Difference [95% CI]", "t", "d_z", "p", "Holm p", "Significant", "VI-SPDAT difference"],
    rows_a, [2.6, 0.45, 1.6, 0.6, 0.55, 0.65, 0.7, 0.85, 1.0],
    "Each row compares Gemini's scores for the same base profiles under two conditions (for example, "
    "the same record labeled American Indian or Alaska Native versus White), in points on the 0–17 "
    "scale; positive values mean the first group scored higher. Paired t-tests; n is the number of "
    "matched pairs. Holm correction across the ten demographic tests (race, gender, and city). "
    "Underdisclosure is shown for reference and is not part of the Holm family. The last column gives "
    "the VI-SPDAT's difference on the same records; * marks a significant result.")
apa_table(
    "B", "Gemini's Scores Relative to the VI-SPDAT, by Race and Gender",
    ["Comparison", "Mean gap (points)", "Difference [95% CI]", "Test statistic", "Effect size", "p", "Holm p", "Significant"],
    rows_b, [2.6, 1.3, 1.6, 1.3, 1.0, 0.65, 0.7, 0.85],
    "The gap is Gemini's score minus the VI-SPDAT's score on the same record. Mean gap gives the "
    "average gap for each group (first group vs. second); the difference compares them, so a positive "
    "value means Gemini raised the first group more relative to the VI-SPDAT. Per-profile means across "
    "the 32 base profiles. Race (all five groups) is a repeated-measures ANOVA with the "
    "Greenhouse–Geisser correction; pairwise comparisons are paired t-tests (df = 31), Holm-adjusted "
    "across the four race contrasts. Because the VI-SPDAT does not use race or gender, these "
    "differences estimate the same quantity as Table A, with profiles rather than records as the unit "
    "of analysis.")

(ROOT / "exports").mkdir(exist_ok=True)
doc.save(ROOT / "exports" / "gemini_tables.docx")
print("Table A"); [print("  " + " | ".join(r)) for r in rows_a]
print("Table B"); [print("  " + " | ".join(r)) for r in rows_b]
print("wrote exports/gemini_tables.docx")
