"""
Compare each AI run's scores with the VI-SPDAT's, by race and by gender.

For every scored instance, gap = AI score - VI-SPDAT score (same record, same disclosure
condition). If an AI run is demographically unbiased, the gap is the same for every race and
both genders; a group whose gap is larger is scored higher by the AI, relative to the
VI-SPDAT, than the reference group.

Tests (per AI run), all paired within base profile so profile difficulty cancels out:
  - Overall gap per group: one-sample t-test of the per-profile mean gap against 0
    (does the AI score this group differently from the VI-SPDAT at all?).
  - Race: repeated-measures ANOVA on the per-profile mean gap across the 5 races, then each
    race vs. White as a paired t-test, Holm-adjusted across the 4 contrasts.
  - Gender: paired t-test of the per-profile mean gap, female vs. male.

Each base profile appears under all 20 conditions (2 genders x 5 races x 2 disclosure) across
the 20 sessions, so every profile contributes to every group.

Writes: data/ai/results/gap_by_group.csv, data/ai/results/gap_tests.csv
        exports/table_gap_by_group.docx (APA 7 table for the paper)
Run:    python scripts/compare_vs_vispdat.py
"""

import csv
import json
from collections import defaultdict
from pathlib import Path

import numpy as np
from scipy import stats

ROOT = Path(__file__).resolve().parents[1]
RES = ROOT / "data" / "ai" / "results"
roster = json.loads((ROOT / "data" / "ai" / "models.json").read_text(encoding="utf-8"))["models"]
RUNS = [(m["label"], m["key"]) for m in roster if (RES / f'{m["key"]}.csv').exists()]

RACES = ["American Indian, Alaska Native, or Indigenous", "Asian or Asian American",
         "Black, African American, or African", "Hispanic/Latino/e/a", "White"]
SHORT = {"American Indian, Alaska Native, or Indigenous": "American Indian or Alaska Native",
         "Asian or Asian American": "Asian", "Black, African American, or African": "Black",
         "Hispanic/Latino/e/a": "Hispanic or Latino", "White": "White"}
GENDERS = ["Male", "Female"]


def holm(ps):
    order = np.argsort(ps)
    adj = [0.0] * len(ps)
    running = 0.0
    for rank, i in enumerate(order):
        running = max(running, min(1.0, (len(ps) - rank) * ps[i]))
        adj[i] = running
    return adj


def paired(a, b):
    d = np.asarray(a) - np.asarray(b)
    t = stats.ttest_1samp(d, 0)
    ci = stats.t.interval(0.95, len(d) - 1, loc=d.mean(), scale=stats.sem(d))
    return {"diff": d.mean(), "lo": ci[0], "hi": ci[1], "t": t.statistic, "df": len(d) - 1,
            "p": t.pvalue, "dz": d.mean() / d.std(ddof=1)}


def rm_anova(matrix):
    """One-way repeated-measures ANOVA; rows = subjects (profiles), columns = levels."""
    X = np.asarray(matrix)
    n, k = X.shape
    grand = X.mean()
    ss_cond = n * ((X.mean(axis=0) - grand) ** 2).sum()
    ss_subj = k * ((X.mean(axis=1) - grand) ** 2).sum()
    ss_total = ((X - grand) ** 2).sum()
    ss_err = ss_total - ss_cond - ss_subj
    df1, df2 = k - 1, (k - 1) * (n - 1)
    F = (ss_cond / df1) / (ss_err / df2)
    eta_p = ss_cond / (ss_cond + ss_err)
    return {"F": F, "df1": df1, "df2": df2, "p": stats.f.sf(F, df1, df2), "eta_p2": eta_p}


group_rows, test_rows, table = [], [], []
for label, key in RUNS:
    rows = list(csv.DictReader((RES / f"{key}.csv").open(encoding="utf-8-sig")))
    # per-profile mean gap and mean scores for each group
    gap = defaultdict(lambda: defaultdict(list))
    ai_s = defaultdict(list)
    vi_s = defaultdict(list)
    for r in rows:
        g = int(r["score"]) - int(r["vispdat_score"])
        for field in ("race", "gender"):
            gap[(field, r[field])][r["profile"]].append(g)
            ai_s[(field, r[field])].append(int(r["score"]))
            vi_s[(field, r[field])].append(int(r["vispdat_score"]))
    profiles = sorted({r["profile"] for r in rows})

    def per_profile(field, level):
        return [np.mean(gap[(field, level)][p]) for p in profiles]

    for field, levels in (("race", RACES), ("gender", GENDERS)):
        for lvl in levels:
            pp = per_profile(field, lvl)
            t = stats.ttest_1samp(pp, 0)
            ci = stats.t.interval(0.95, len(pp) - 1, loc=np.mean(pp), scale=stats.sem(pp))
            group_rows.append({
                "run": label, "factor": field, "group": SHORT.get(lvl, lvl), "n": len(ai_s[(field, lvl)]),
                "mean_ai": np.mean(ai_s[(field, lvl)]), "mean_vispdat": np.mean(vi_s[(field, lvl)]),
                "gap": np.mean(pp), "gap_ci_low": ci[0], "gap_ci_high": ci[1], "t_vs_0": t.statistic, "p_vs_0": t.pvalue,
            })

    # race: omnibus, then each race vs. White
    mat = np.column_stack([per_profile("race", r) for r in RACES])
    om = rm_anova(mat)
    contrasts = [paired(per_profile("race", r), per_profile("race", "White")) for r in RACES[:-1]]
    padj = holm([c["p"] for c in contrasts])
    test_rows.append({"run": label, "test": "Race (omnibus, RM-ANOVA)", "estimate": "",
                      "stat": f'F({om["df1"]}, {om["df2"]}) = {om["F"]:.2f}', "p": om["p"], "p_holm": "",
                      "effect": f'partial eta^2 = {om["eta_p2"]:.3f}'})
    for r, c, pa in zip(RACES[:-1], contrasts, padj):
        test_rows.append({"run": label, "test": f"{SHORT[r]} vs. White", "estimate": c["diff"],
                          "stat": f't({c["df"]}) = {c["t"]:.2f}', "p": c["p"], "p_holm": pa,
                          "effect": f'd_z = {c["dz"]:.2f}', "lo": c["lo"], "hi": c["hi"]})
    g = paired(per_profile("gender", "Female"), per_profile("gender", "Male"))
    test_rows.append({"run": label, "test": "Female vs. male", "estimate": g["diff"],
                      "stat": f't({g["df"]}) = {g["t"]:.2f}', "p": g["p"], "p_holm": g["p"],
                      "effect": f'd_z = {g["dz"]:.2f}', "lo": g["lo"], "hi": g["hi"]})
    table.append((label, om, contrasts, padj, g))

RES.mkdir(parents=True, exist_ok=True)
with (RES / "gap_by_group.csv").open("w", encoding="utf-8-sig", newline="") as fh:
    w = csv.DictWriter(fh, fieldnames=list(group_rows[0]))
    w.writeheader()
    for r in group_rows:
        w.writerow({k: (round(v, 4) if isinstance(v, float) else v) for k, v in r.items()})
fields = ["run", "test", "estimate", "lo", "hi", "stat", "p", "p_holm", "effect"]
with (RES / "gap_tests.csv").open("w", encoding="utf-8-sig", newline="") as fh:
    w = csv.DictWriter(fh, fieldnames=fields)
    w.writeheader()
    for r in test_rows:
        w.writerow({k: (round(r.get(k), 4) if isinstance(r.get(k), float) else r.get(k, "")) for k in fields})

# ---------- console summary ----------
for label, om, contrasts, padj, g in table:
    print(f"\n{label}")
    print(f"  race omnibus: F({om['df1']},{om['df2']}) = {om['F']:.2f}, p = {om['p']:.3f}, partial eta2 = {om['eta_p2']:.3f}")
    for r, c, pa in zip(RACES[:-1], contrasts, padj):
        print(f"  {SHORT[r]:32} vs White: {c['diff']:+.2f} [{c['lo']:+.2f}, {c['hi']:+.2f}]  p = {c['p']:.3f}, Holm p = {pa:.3f}")
    print(f"  {'Female':32} vs male : {g['diff']:+.2f} [{g['lo']:+.2f}, {g['hi']:+.2f}]  p = {g['p']:.3f}")
print("\nMean gap (AI - VI-SPDAT) by group:")
for r in group_rows:
    print(f"  {r['run']:24} {r['group']:32} AI {r['mean_ai']:.2f}  VI {r['mean_vispdat']:.2f}  gap {r['gap']:+.2f} [{r['gap_ci_low']:+.2f}, {r['gap_ci_high']:+.2f}]")


# ---------- APA 7 table for the paper ----------
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def gg_p(matrix):
    """Repeated-measures ANOVA p-value with the Greenhouse-Geisser sphericity correction."""
    X = np.asarray(matrix)
    n, k = X.shape
    a = rm_anova(X)
    C = np.eye(k) - 1 / k
    S = C @ np.cov(X, rowvar=False) @ C
    eps = np.trace(S) ** 2 / ((k - 1) * np.sum(S * S))
    return stats.f.sf(a["F"], a["df1"] * eps, a["df2"] * eps), eps


def rule(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:color"), "000000")
    borders.append(el)


def fp(p):
    return "< .001" if p < 0.001 else f"{p:.3f}".lstrip("0")


doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Inches(1))
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(12)

p = doc.add_paragraph()
p.add_run("Table 2").bold = True
p.paragraph_format.space_after = Pt(0)
p = doc.add_paragraph()
p.add_run("Difference Between AI and VI-SPDAT Scores by Race and Gender").italic = True
p.paragraph_format.space_after = Pt(6)

labels = [r[0] for r in RUNS]
t = doc.add_table(rows=1, cols=1 + len(labels))
t.alignment = WD_TABLE_ALIGNMENT.CENTER
head = ["Group"] + labels
for i, h in enumerate(head):
    c = t.rows[0].cells[i]
    c.text = ""
    c.paragraphs[0].add_run(h).font.size = Pt(10)
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
    rule(c, "top", 8)
    rule(c, "bottom", 4)


def add_row(cells, italic=False, last=False):
    row = t.add_row().cells
    for i, v in enumerate(cells):
        row[i].text = ""
        run = row[i].paragraphs[0].add_run(v)
        run.font.size = Pt(10)
        run.italic = italic
        row[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
        if last:
            rule(row[i], "bottom", 8)


by = {(r["run"], r["factor"], r["group"]): r for r in group_rows}
add_row(["Race (AI minus VI-SPDAT, points)"] + [""] * len(labels), italic=True)
for lvl in RACES:
    g = SHORT[lvl]
    add_row([f"   {g}"] + [f'{by[(l, "race", g)]["gap"]:.2f} [{by[(l, "race", g)]["gap_ci_low"]:.2f}, {by[(l, "race", g)]["gap_ci_high"]:.2f}]' for l in labels])
omni = []
for label, key in RUNS:
    rows = list(csv.DictReader((RES / f"{key}.csv").open(encoding="utf-8-sig")))
    gg = defaultdict(list)
    for r in rows:
        gg[(r["profile"], r["race"])].append(int(r["score"]) - int(r["vispdat_score"]))
    P = sorted({r["profile"] for r in rows})
    X = np.array([[np.mean(gg[(pr, rc)]) for rc in RACES] for pr in P])
    a = rm_anova(X)
    pgg, eps = gg_p(X)
    omni.append((a, pgg))
add_row(["   Race difference, F (p)"] + [f'{a["F"]:.2f} ({fp(pgg)})' for a, pgg in omni])
add_row(["Gender (AI minus VI-SPDAT, points)"] + [""] * len(labels), italic=True)
for g in GENDERS:
    add_row([f"   {g}"] + [f'{by[(l, "gender", g)]["gap"]:.2f} [{by[(l, "gender", g)]["gap_ci_low"]:.2f}, {by[(l, "gender", g)]["gap_ci_high"]:.2f}]' for l in labels])
gend = {r["run"]: r for r in test_rows if r["test"] == "Female vs. male"}
add_row(["   Female minus male, t (p)"] + [f'{float(gend[l]["stat"].split("= ")[1]):.2f} ({fp(gend[l]["p"])})' for l in labels], last=True)

for row in t.rows:
    row.cells[0].width = Inches(2.6)
    for c in row.cells[1:]:
        c.width = Inches(1.6)

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(6)
note.add_run("Note. ").italic = True
note.add_run(
    "Values are mean differences (AI score minus VI-SPDAT score on the same record) with 95% "
    "confidence intervals, computed within each of the 32 base profiles and averaged. Each profile "
    "was scored under every race, gender, and disclosure condition, so profile differences cancel. "
    "A demographically unbiased AI shows the same difference for every group. Race differences were "
    "tested with a repeated-measures ANOVA, F(4, 124), with the Greenhouse–Geisser correction; "
    "gender with a paired t-test, t(31). No race or gender difference was significant after Holm "
    "correction across the four AI runs except the female–male difference for ChatGPT (default "
    "reasoning). The VI-SPDAT does not use race or gender, so its scores are identical across groups."
).font.size = Pt(10)
(ROOT / "exports").mkdir(exist_ok=True)
doc.save(ROOT / "exports" / "table_gap_by_group.docx")
print("\nwrote exports/table_gap_by_group.docx")


# ---------- Table 3: every significance test, with Holm correction across the AI runs ----------
tests = []  # (run, comparison, estimate, ci, stat, p)
for (label, key), (a, pgg) in zip(RUNS, omni):
    tests.append((label, "Race (all five groups)", None, None, f'F(4, 124) = {a["F"]:.2f}', pgg, "race_omni"))
for r in test_rows:
    if r["test"] == "Race (omnibus, RM-ANOVA)":
        continue
    fam = "gender" if r["test"] == "Female vs. male" else "race_pair"
    tests.append((r["run"], r["test"], r["estimate"], (r["lo"], r["hi"]), r["stat"], r["p"], fam))

# Holm within each family of tests: race omnibus across runs, race pairs across runs x 4, gender across runs
adj = {}
for fam in ("race_omni", "race_pair", "gender"):
    idx = [i for i, x in enumerate(tests) if x[6] == fam]
    for i, pa in zip(idx, holm([tests[i][5] for i in idx])):
        adj[i] = pa

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(18)
p.add_run("Table 3").bold = True
p.paragraph_format.space_after = Pt(0)
p = doc.add_paragraph()
p.add_run("Significance Tests for Race and Gender Differences in the AI–VI-SPDAT Score Gap").italic = True
p.paragraph_format.space_after = Pt(6)

heads = ["AI run", "Comparison", "Difference [95% CI]", "Test statistic", "p", "Holm p", "Significant"]
t = doc.add_table(rows=1, cols=len(heads))
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(heads):
    c = t.rows[0].cells[i]
    c.text = ""
    c.paragraphs[0].add_run(h).font.size = Pt(10)
    rule(c, "top", 8)
    rule(c, "bottom", 4)
order = sorted(range(len(tests)), key=lambda i: ([r[0] for r in RUNS].index(tests[i][0]),
                                                  {"race_omni": 0, "race_pair": 1, "gender": 2}[tests[i][6]]))
prev = None
for n, i in enumerate(order):
    run, comp, est, ci, stat, pv, fam = tests[i]
    cells = t.add_row().cells
    vals = [run if run != prev else "", comp,
            "" if est is None else f"{est:+.2f} [{ci[0]:+.2f}, {ci[1]:+.2f}]",
            stat.replace("t(", "t(").replace("F(", "F("), fp(pv), fp(adj[i]), "Yes" if adj[i] < 0.05 else "No"]
    prev = run
    for j, v in enumerate(vals):
        cells[j].text = ""
        run_ = cells[j].paragraphs[0].add_run(v)
        run_.font.size = Pt(10)
        if j == 6 and v == "Yes":
            run_.bold = True
        if n == len(order) - 1:
            rule(cells[j], "bottom", 8)
widths = [1.6, 2.6, 1.6, 1.3, 0.6, 0.7, 0.8]
for row in t.rows:
    for c, w in zip(row.cells, widths):
        c.width = Inches(w)

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(6)
note.add_run("Note. ").italic = True
note.add_run(
    "Differences are in points on the 0–17 scale and compare the AI–VI-SPDAT gap between groups "
    "(for example, the gap for American Indian or Alaska Native cases minus the gap for White cases); "
    "positive values mean the AI scores the first group higher relative to the VI-SPDAT. Race (all five "
    "groups) is a repeated-measures ANOVA with the Greenhouse–Geisser correction; pairwise and gender "
    "comparisons are paired t-tests across the 32 base profiles, df = 31. Holm-adjusted p values correct "
    "within each family of tests across all four AI runs (race omnibus, 4 tests; race vs. White, 16 tests; "
    "female vs. male, 4 tests). Significant: Holm p < .05."
).font.size = Pt(10)
doc.save(ROOT / "exports" / "table_gap_by_group.docx")
print("added Table 3 to exports/table_gap_by_group.docx")
