"""
Equal-capacity housing allocation: VI-SPDAT vs. an AI run, for every characteristic.

Both tools get the same number of housing places: as many as there are truly high-need records
(true score 8 or more), 280 of 640. Each tool fills them with its highest-scoring records; ties at
the cutoff share the remaining places evenly. With the same capacity, each group's share of the
places can be compared directly between tools.

Fair share:
  race, gender, city, disclosure   the group's share of the population (every group is the same
                                   size and has the same real need by design)
  profile characteristics          the group's share of everyone who truly needs housing

Deviation = share of places - fair share (percentage points). Positive: the group gets more than its
fair share.

Tests (each tool): permutation, 5,000 shuffles. Demographic labels are shuffled within each base
profile; profile characteristics are reassigned across the 32 base profiles. Statistic: sum of
squared deviations across the characteristic's groups. Holm correction across profile
characteristics.

Run:    python scripts/equal_capacity.py --run gemini
Writes: data/ai/results/equal_capacity_<run>.csv, exports/equal_capacity_<run>.docx,
        exports/figures/equal_capacity_{demographics,traits}_<run>.png
"""

import argparse
import csv
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
RES = ROOT / "data" / "ai" / "results"
OUT = ROOT / "exports"
(OUT / "figures").mkdir(parents=True, exist_ok=True)
ap = argparse.ArgumentParser()
ap.add_argument("--run", default="gemini")
ap.add_argument("--n", type=int, default=5000)
ap.add_argument("--out", help="output file name (default: equal_capacity_<run>.docx)")
args = ap.parse_args()
roster = json.loads((ROOT / "data" / "ai" / "models.json").read_text(encoding="utf-8"))["models"]
AI = next(m["label"] for m in roster if m["key"] == args.run)
profiles = {p["id"]: p for p in json.loads((ROOT / "data" / "base_profiles.json").read_text(encoding="utf-8"))["profiles"]}
inst = json.loads((ROOT / "data" / "vispdat_instrument.json").read_text(encoding="utf-8"))

df = pd.DataFrame(list(csv.DictReader((RES / f"{args.run}.csv").open(encoding="utf-8-sig"))))
for c in ("score", "vispdat_score", "true_score"):
    df[c] = df[c].astype(int)
truth = (df.true_score >= 8).values.astype(float)
K = int(truth.sum())


def places(scores):
    """Expected places each record receives when the top K scores are housed; ties split evenly."""
    s = np.asarray(scores)
    cut = np.sort(s)[::-1][K - 1]
    w = (s > cut).astype(float)
    at = s == cut
    w[at] = (K - (s > cut).sum()) / at.sum()
    return w


TOOLS = {"VI-SPDAT": places(df.vispdat_score.values), AI: places(df.score.values)}
SHORT = {"American Indian, Alaska Native, or Indigenous": "American Indian or Alaska Native",
         "Asian or Asian American": "Asian", "Black, African American, or African": "Black",
         "Hispanic/Latino/e/a": "Hispanic or Latino"}
df["race"] = df["race"].map(lambda r: SHORT.get(r, r))
df["location"] = df["location"].str.replace(r", [A-Z][a-z]+( [A-Z][a-z]+)?$", "", regex=True)
prof = df.profile.values
prof_ids = np.array(sorted(df.profile.unique()))
idx_by_prof = [np.where(prof == p)[0] for p in prof_ids]
rng = np.random.default_rng(20261011)
NAMES = {"d1_physical": "Physical illness", "c1_money": "Money problems (finances)", "a2_duration": "Long homelessness",
         "a1_sleep": "Unsheltered or unstable housing", "pre_age": "Age 60 or older", "c2_activity": "No meaningful daily activity",
         "c3_selfcare": "Self-care problems", "c4_relationships": "Social relationship problems", "b1_emergency": "Emergency service use",
         "b2_harm": "Risk of harm", "b3_legal": "Legal issues", "b4_exploitation": "Risk of exploitation",
         "d2_substance": "Substance use", "d3_mental": "Mental health", "d5_medications": "Medication problems",
         "d6_trauma": "Abuse and trauma"}


def deviations(labels, w, fair_by):
    """Share of places minus fair share, per group."""
    share = pd.Series(w).groupby(labels).sum() / w.sum()
    fair = fair_by(labels)
    return share - fair


def fair_population(labels):
    return pd.Series(1.0, index=range(len(labels))).groupby(labels).sum() / len(labels)


def fair_need(labels):
    return pd.Series(truth).groupby(labels).sum() / truth.sum()


def stat(labels, w, fair_by):
    return float((deviations(labels, w, fair_by) ** 2).sum())


def perm_within(labels, w):
    obs = stat(labels, w, fair_population)
    hits = 0
    for _ in range(args.n):
        lab = labels.copy()
        for i in idx_by_prof:
            lab[i] = labels[i][rng.permutation(len(i))]
        hits += stat(lab, w, fair_population) >= obs - 1e-12
    return (hits + 1) / (args.n + 1)


def perm_across(with_p, w):
    obs = stat(np.isin(prof, with_p), w, fair_need)
    hits = 0
    for _ in range(args.n):
        pick = rng.choice(prof_ids, size=len(with_p), replace=False)
        hits += stat(np.isin(prof, pick), w, fair_need) >= obs - 1e-12
    return (hits + 1) / (args.n + 1)


def holm(ps):
    order = np.argsort(ps)
    adj, run = [0.0] * len(ps), 0.0
    for rank, i in enumerate(order):
        run = max(run, min(1.0, (len(ps) - rank) * ps[i]))
        adj[i] = run
    return adj


rows, tests = [], []
for field, title in (("gender", "Gender"), ("race", "Race"), ("location", "City"), ("disclosure", "Disclosure")):
    labels = df[field].values
    t = {"section": "Demographic", "characteristic": title}
    for tool, w in TOOLS.items():
        t[f"p_{tool}"] = 1.0 if (tool == "VI-SPDAT" and field != "disclosure") else perm_within(labels.copy(), w)
    tests.append(t)
    fair = fair_population(labels)
    for lvl in sorted(df[field].unique()):
        rows.append({"section": "Demographic", "characteristic": title, "group": lvl, "fair": fair[lvl],
                     **{tool: (w[labels == lvl].sum() / w.sum()) for tool, w in TOOLS.items()}})

seen, trait_tests = {}, []
for d in inst["domains"]:
    for ind in d["indicators"]:
        if ind["id"] not in NAMES:
            continue
        with_p = [p for p in prof_ids if ind["id"] in profiles[p]["indicators"]]
        key = frozenset(with_p)
        if key in seen:
            continue
        seen[key] = NAMES[ind["id"]]
        has = np.isin(prof, with_p)
        t = {"section": "Profile", "characteristic": NAMES[ind["id"]]}
        for tool, w in TOOLS.items():
            t[f"p_{tool}"] = perm_across(with_p, w)
        trait_tests.append(t)
        for g, m in (("With", has), ("Without", ~has)):
            rows.append({"section": "Profile", "characteristic": NAMES[ind["id"]], "group": g,
                         "fair": truth[m].sum() / truth.sum(),
                         **{tool: (w[m].sum() / w.sum()) for tool, w in TOOLS.items()}})
for tool in TOOLS:
    for t, pa in zip(trait_tests, holm([t[f"p_{tool}"] for t in trait_tests])):
        t[f"p_{tool}"] = pa
tests = pd.DataFrame(tests + trait_tests)
table = pd.DataFrame(rows).merge(tests, on=["section", "characteristic"])
for tool in TOOLS:
    table[f"dev_{tool}"] = table[tool] - table["fair"]
table.to_csv(RES / f"equal_capacity_{args.run}.csv", index=False, encoding="utf-8-sig")


def fp(p):
    return "< .001" if p < 0.001 else f"{p:.3f}".lstrip("0")


# ---------------- figures: deviation from fair share ----------------
plt.rcParams.update({"font.family": "Times New Roman", "font.size": 10, "axes.spines.top": False,
                     "axes.spines.right": False, "savefig.dpi": 300, "savefig.bbox": "tight"})


def draw(sub, path, group_label):
    sub = sub.reset_index(drop=True)
    labels = [group_label(r) for _, r in sub.iterrows()]
    y = np.arange(len(sub))[::-1]
    fig, ax = plt.subplots(figsize=(7.0, 0.32 * len(sub) + 1.4))
    ax.axvline(0, color="black", linewidth=1)
    ax.barh(y + 0.18, 100 * sub["dev_VI-SPDAT"], height=0.34, color="black", label="VI-SPDAT")
    ax.barh(y - 0.18, 100 * sub[f"dev_{AI}"], height=0.34, color="0.65", edgecolor="black", linewidth=0.5, label=AI)
    # separators and significance marks between characteristics
    prev = None
    for yi, (_, r) in zip(y, sub.iterrows()):
        if r.characteristic != prev and prev is not None:
            ax.axhline(yi + 0.5, color="0.85", linewidth=0.6)
        if r.characteristic != prev:
            marks = "".join(["†" if r["p_VI-SPDAT"] < 0.05 else "", "*" if r[f"p_{AI}"] < 0.05 else ""])
            if marks:
                ax.text(1.01, yi, marks, transform=ax.get_yaxis_transform(), va="center", fontsize=11)
        prev = r.characteristic
    ax.set_yticks(y)
    ax.set_yticklabels(labels, fontsize=9)
    ax.tick_params(axis="y", length=0)
    lim = max(4, np.ceil(100 * np.abs(sub[["dev_VI-SPDAT", f"dev_{AI}"]].values).max() + 1))
    ax.set_xlim(-lim, lim)
    ax.set_xlabel("Share of housing places minus fair share (percentage points)")
    fig.legend(*ax.get_legend_handles_labels(), loc="lower center", ncol=2, frameon=False, bbox_to_anchor=(0.5, -0.02))
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


f1 = OUT / "figures" / f"equal_capacity_demographics_{args.run}.png"
f2 = OUT / "figures" / f"equal_capacity_traits_{args.run}.png"
draw(table[table.section == "Demographic"], f1, lambda r: f"{r.characteristic}: {r.group}")
draw(table[(table.section == "Profile") & (table.group == "With")], f2, lambda r: r.characteristic)

# ---------------- Word file ----------------
doc = Document()
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(doc.sections[0], side, Inches(1))
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(12)


def rule(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:color"), "000000")
    b.append(el)


def heading(kind, n, title):
    p = doc.add_paragraph()
    p.add_run(f"{kind} {n}").bold = True
    p.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph()
    p.add_run(title).italic = True
    p.paragraph_format.space_after = Pt(6)


def note(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.add_run("Note. ").italic = True
    p.add_run(text).font.size = Pt(10)


r = doc.add_paragraph().add_run(f"Equal-Capacity Housing Allocation: VI-SPDAT and {AI}")
r.bold = True
r.font.size = Pt(16)
common = (f"Both tools fill the same {K} housing places (the number of records that truly need housing, true score 8 "
          "or more) with their highest-scoring records; ties at the cutoff share places evenly. Bars show each group's "
          "share of the places minus its fair share.")
heading("Figure", 1, "Share of Housing Places Relative to Fair Share, by Demographic Group")
doc.add_picture(str(f1), width=Inches(6.5))
note(common + " Fair share = the group's share of the population. * " + AI + " and † VI-SPDAT differ from fair shares "
     "at p < .05 (permutation test, labels shuffled within each base profile, 5,000 shuffles). The VI-SPDAT does not "
     "use race, gender, or city.")
doc.add_page_break()
heading("Figure", 2, "Share of Housing Places Going to People With Each Characteristic, Relative to Fair Share")
doc.add_picture(str(f2), width=Inches(6.5))
note(common + " Fair share = the group's share of everyone who truly needs housing. People without the characteristic "
     "receive the mirror image. * " + AI + " and † VI-SPDAT differ from fair shares at Holm-corrected p < .05 "
     "(permutation test reassigning the characteristic across the 32 base profiles, 5,000 reassignments).")
doc.add_page_break()

heading("Table", 1, "Share of Housing Places Going to Each Group Under Equal Capacity")
heads = ["Characteristic", "Group", "Fair share", "VI-SPDAT share", f"{AI} share", "VI-SPDAT difference",
         f"{AI} difference", "VI-SPDAT p", f"{AI} p"]
t = doc.add_table(rows=1, cols=len(heads))
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(heads):
    c = t.rows[0].cells[i]
    c.text = ""
    c.paragraphs[0].add_run(h).font.size = Pt(8.5)
    rule(c, "top", 8)
    rule(c, "bottom", 4)
prev = None
for n, (_, r) in enumerate(table.iterrows()):
    first = r.characteristic != prev
    vis_p = "—" if (r.section == "Demographic" and r.characteristic != "Disclosure") else fp(r["p_VI-SPDAT"])
    vals = [r.characteristic if first else "", r.group, f"{100 * r.fair:.1f}%", f"{100 * r['VI-SPDAT']:.1f}%",
            f"{100 * r[AI]:.1f}%", f"{100 * r['dev_VI-SPDAT']:+.1f}", f"{100 * r[f'dev_{AI}']:+.1f}",
            vis_p if first else "", fp(r[f"p_{AI}"]) if first else ""]
    prev = r.characteristic
    cells = t.add_row().cells
    for i, v in enumerate(vals):
        cells[i].text = ""
        cells[i].paragraphs[0].add_run(v.replace("-", "−") if i in (5, 6) else v).font.size = Pt(8.5)
        cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if i < 2 else WD_ALIGN_PARAGRAPH.CENTER
        if n == len(table) - 1:
            rule(cells[i], "bottom", 8)
for row in t.rows:
    for c, w in zip(row.cells, [1.4, 1.2, 0.6, 0.65, 0.6, 0.65, 0.6, 0.55, 0.55]):
        c.width = Inches(w)
note(common + " Difference = share minus fair share, in percentage points. Fair share: population share for race, "
     "gender, city, and disclosure; share of everyone who truly needs housing for profile characteristics. p: "
     "permutation tests as in Figures 1 and 2 (profile characteristics Holm-corrected). — : the VI-SPDAT does not "
     "use race, gender, or city, so its shares equal the fair shares exactly.")
out_path = OUT / (args.out or f"equal_capacity_{args.run}.docx")
doc.save(out_path)
print(f"wrote {out_path.relative_to(ROOT)}")
for _, r in table.iterrows():
    print(f"{r.characteristic[:30]:30} {r.group[:32]:32} fair {100*r.fair:5.1f}  VI {100*r['VI-SPDAT']:5.1f} ({100*r['dev_VI-SPDAT']:+5.1f})  "
          f"{AI} {100*r[AI]:5.1f} ({100*r[f'dev_{AI}']:+5.1f})  pVI {r['p_VI-SPDAT']:.3f} pAI {r[f'p_{AI}']:.3f}")
