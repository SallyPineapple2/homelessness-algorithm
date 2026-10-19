"""
Allocation comparison for every characteristic: what share of each group truly needs housing, and
what share each tool gives housing to.

"Housing" = a referral to Permanent Supportive Housing (score 8 or more). "Truly needs it" = the
record's true score is 8 or more. For a fair tool, the share referred in each group matches the
share that truly needs it; a tendency shows up as one group being referred above (or below) its
real need by more than the other group.

Groups: race, gender, city, disclosure (assigned by design) and the 16 VI-SPDAT characteristics
written into the base profiles (with vs. without each).

Test (for each characteristic, each tool): is (referred - truly needs) different between the
groups? Permutation test, 5,000 shuffles: demographic labels are shuffled within each base profile;
profile characteristics are reassigned across the 32 base profiles. Holm correction within the
profile characteristics.

Run:    python scripts/allocation_comparison.py --run gemini
Writes: data/ai/results/allocation_comparison_<run>.csv, exports/allocation_comparison_<run>.docx
"""

import argparse
import csv
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
RES = ROOT / "data" / "ai" / "results"
OUT = ROOT / "exports"
ap = argparse.ArgumentParser()
ap.add_argument("--run", default="gemini")
ap.add_argument("--n", type=int, default=5000)
ap.add_argument("--out", help="output file name (default: allocation_comparison_<run>.docx)")
args = ap.parse_args()
roster = json.loads((ROOT / "data" / "ai" / "models.json").read_text(encoding="utf-8"))["models"]
AI = next(m["label"] for m in roster if m["key"] == args.run)
profiles = {p["id"]: p for p in json.loads((ROOT / "data" / "base_profiles.json").read_text(encoding="utf-8"))["profiles"]}
inst = json.loads((ROOT / "data" / "vispdat_instrument.json").read_text(encoding="utf-8"))

df = pd.DataFrame(list(csv.DictReader((RES / f"{args.run}.csv").open(encoding="utf-8-sig"))))
for c in ("score", "vispdat_score", "true_score"):
    df[c] = df[c].astype(int)
truth = (df.true_score >= 8).values.astype(float)
TOOLS = {"VI-SPDAT": (df.vispdat_score >= 8).values.astype(float), AI: (df.score >= 8).values.astype(float)}
SHORT = {"American Indian, Alaska Native, or Indigenous": "American Indian or Alaska Native",
         "Asian or Asian American": "Asian", "Black, African American, or African": "Black",
         "Hispanic/Latino/e/a": "Hispanic or Latino"}
df["race"] = df["race"].map(lambda r: SHORT.get(r, r))
df["location"] = df["location"].str.replace(r", [A-Z][a-z]+( [A-Z][a-z]+)?$", "", regex=True)
prof = df.profile.values
prof_ids = np.array(sorted(df.profile.unique()))
idx_by_prof = [np.where(prof == p)[0] for p in prof_ids]
rng = np.random.default_rng(20261011)
TRAIT_NAMES = {"d1_physical": "Physical health (illness)", "c1_money": "Money management (finances)",
               "a2_duration": "Long homelessness", "a1_sleep": "Unsheltered or unstable housing",
               "pre_age": "Age 60 or older"}


def shares_of_whole(m):
    """This group's share of the population, of everyone truly needing housing, and of each tool's housing."""
    return {"pop_share": m.mean(), "need_share": truth[m].sum() / truth.sum(),
            "vis_share": TOOLS["VI-SPDAT"][m].sum() / TOOLS["VI-SPDAT"].sum(),
            "ai_share": TOOLS[AI][m].sum() / TOOLS[AI].sum()}


def excess_spread(labels, refer):
    """Spread across groups of (share referred - share truly needing it)."""
    s = pd.DataFrame({"g": labels, "r": refer, "t": truth}).groupby("g").mean()
    e = s.r - s.t
    return (e.max() - e.min()) if len(e) > 1 else 0.0


def perm_within(labels, refer):
    obs = excess_spread(labels, refer)
    hits = 0
    for _ in range(args.n):
        lab = labels.copy()
        for i in idx_by_prof:
            lab[i] = labels[i][rng.permutation(len(i))]
        hits += excess_spread(lab, refer) >= obs - 1e-12
    return (hits + 1) / (args.n + 1)


def perm_across(with_profiles, refer):
    has = np.isin(prof, with_profiles)
    obs = excess_spread(has, refer)
    hits = 0
    for _ in range(args.n):
        p = rng.choice(prof_ids, size=len(with_profiles), replace=False)
        hits += excess_spread(np.isin(prof, p), refer) >= obs - 1e-12
    return (hits + 1) / (args.n + 1)


def holm(ps):
    order = np.argsort(ps)
    adj, run = [0.0] * len(ps), 0.0
    for rank, i in enumerate(order):
        run = max(run, min(1.0, (len(ps) - rank) * ps[i]))
        adj[i] = run
    return adj


rows, tests = [], []
# demographic groups (assigned by design)
for field, title in (("gender", "Gender"), ("race", "Race"), ("location", "City"), ("disclosure", "Disclosure")):
    labels = df[field].values
    ps = {t: (1.0 if t == "VI-SPDAT" and field != "disclosure" else perm_within(labels.copy(), r)) for t, r in TOOLS.items()}
    for lvl in sorted(df[field].unique()):
        m = labels == lvl
        rows.append({"section": "Demographic", "characteristic": title, "group": lvl, "records": int(m.sum()),
                     "truly": truth[m].mean(), "VI-SPDAT": TOOLS["VI-SPDAT"][m].mean(), AI: TOOLS[AI][m].mean(),
                     **shares_of_whole(m)})
    tests.append({"section": "Demographic", "characteristic": title, "p_vis": ps["VI-SPDAT"], "p_ai": ps[AI]})

# profile characteristics
seen = {}
trait_tests = []
for d in inst["domains"]:
    for ind in d["indicators"]:
        with_p = [p for p in prof_ids if ind["id"] in profiles[p]["indicators"]]
        if not with_p or len(with_p) == len(prof_ids):
            continue
        name = TRAIT_NAMES.get(ind["id"], ind["name"])
        key = frozenset(with_p)
        if key in seen:
            for r in rows:
                if r["characteristic"] == seen[key]:
                    r["characteristic"] = f"{seen[key]} / {name}"
            for t in trait_tests:
                if t["characteristic"] == seen[key]:
                    t["characteristic"] = f"{seen[key]} / {name}"
            seen[key] = f"{seen[key]} / {name}"
            continue
        seen[key] = name
        has = np.isin(prof, with_p)
        for g, m in (("With", has), ("Without", ~has)):
            rows.append({"section": "Profile", "characteristic": name, "group": g, "records": int(m.sum()),
                         "truly": truth[m].mean(), "VI-SPDAT": TOOLS["VI-SPDAT"][m].mean(), AI: TOOLS[AI][m].mean(),
                         **shares_of_whole(m)})
        trait_tests.append({"section": "Profile", "characteristic": name,
                            "p_vis": perm_across(with_p, TOOLS["VI-SPDAT"]), "p_ai": perm_across(with_p, TOOLS[AI])})
for key in ("p_vis", "p_ai"):
    for t, pa in zip(trait_tests, holm([t[key] for t in trait_tests])):
        t[key + "_holm"] = pa
tests += trait_tests
for t in tests:
    t.setdefault("p_vis_holm", t["p_vis"])
    t.setdefault("p_ai_holm", t["p_ai"])

table = pd.DataFrame(rows)
tests = pd.DataFrame(tests)
table.merge(tests, on=["section", "characteristic"]).to_csv(RES / f"allocation_comparison_{args.run}.csv",
                                                             index=False, encoding="utf-8-sig")


# ---------------- figures ----------------
plt.rcParams.update({"font.family": "Times New Roman", "font.size": 10, "axes.spines.top": False,
                     "axes.spines.right": False, "savefig.dpi": 300, "savefig.bbox": "tight"})
FILL = [("truly", "Truly needs housing", "white"), ("VI-SPDAT", "VI-SPDAT gives housing", "black"),
        (AI, f"{AI} gives housing", "0.6")]


def fp(p):
    return "< .001" if p < 0.001 else "= " + f"{p:.3f}".lstrip("0")


def draw(sub, path, title_for):
    chars = list(dict.fromkeys(sub.characteristic))
    heights = [len(sub[sub.characteristic == c]) for c in chars]
    fig, axes = plt.subplots(len(chars), 1, figsize=(7.0, 0.55 * sum(heights) + 0.62 * len(chars) + 0.6),
                             gridspec_kw={"height_ratios": heights}, sharex=True)
    axes = np.atleast_1d(axes)
    for ax, c in zip(axes, chars):
        s = sub[sub.characteristic == c].reset_index(drop=True)
        y = np.arange(len(s))[::-1]
        for (col, name, face), off in zip(FILL, (-0.26, 0, 0.26)):
            ax.barh(y - off, 100 * s[col], height=0.24, color=face, edgecolor="black", linewidth=0.6, label=name)
            for yi, v in zip(y - off, 100 * s[col]):
                ax.text(v + 1, yi, f"{v:.0f}%", va="center", fontsize=7.5)
        ax.set_yticks(y)
        ax.set_yticklabels(s.group)
        ax.tick_params(axis="y", length=0)
        ax.set_xlim(0, 112)
        t = tests[tests.characteristic == c].iloc[0]
        ax.set_title(title_for(c, t), loc="left", fontsize=9.5)
    axes[-1].set_xlabel("Share of the group (%)")
    fig.legend(*axes[0].get_legend_handles_labels(), loc="lower center", ncol=3, frameon=False,
               bbox_to_anchor=(0.5, -0.03 if len(chars) > 5 else -0.07))
    fig.tight_layout(h_pad=0.9)
    fig.savefig(path)
    plt.close(fig)


f1 = OUT / "figures" / f"allocation_comparison_demographics_{args.run}.png"
f2 = OUT / "figures" / f"allocation_comparison_traits1_{args.run}.png"
f3 = OUT / "figures" / f"allocation_comparison_traits2_{args.run}.png"
(OUT / "figures").mkdir(parents=True, exist_ok=True)
draw(table[table.section == "Demographic"], f1,
     lambda c, t: f"{c} — {AI}: p {fp(t.p_ai)}" + (f"; VI-SPDAT: p {fp(t.p_vis)}" if c == "Disclosure" else ""))
prof_chars = list(dict.fromkeys(table[table.section == "Profile"].characteristic))
half = (len(prof_chars) + 1) // 2
draw(table[table.characteristic.isin(prof_chars[:half])], f2, lambda c, t: c)
draw(table[table.characteristic.isin(prof_chars[half:])], f3, lambda c, t: c)

# ---------------- Word file ----------------
doc = Document()
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(doc.sections[0], side, Inches(1))
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(12)


def rule(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:color"), "000000")
    b.append(el)


def heading(kind, n, title):
    p = doc.add_paragraph()
    p.add_run(f"{kind} {n}").bold = True
    p.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph()
    p.add_run(title).italic = True
    p.paragraph_format.space_after = Pt(6)


def note(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.add_run("Note. ").italic = True
    p.add_run(text).font.size = Pt(10)


p = doc.add_paragraph()
r = p.add_run(f"Housing Allocation by Group: VI-SPDAT and {AI}")
r.bold = True
r.font.size = Pt(16)

heading("Figure", 1, "Share of Each Demographic Group That Truly Needs Housing and That Each Tool Gives Housing To")
doc.add_picture(str(f1), width=Inches(6.5))
note("Housing = a referral to Permanent Supportive Housing (score 8 or more). Truly needs housing = true score "
     "8 or more. p: permutation test of whether the gap between share given housing and share truly needing it "
     "differs between groups (labels shuffled within each base profile, 5,000 shuffles). The VI-SPDAT does not use "
     "race, gender, or city.")
doc.add_page_break()
for n_, f_, part in ((2, f2, "Part 1"), (3, f3, "Part 2")):
    heading("Figure", n_, "Share of People With and Without Each Characteristic That Truly Need Housing and That "
                          f"Each Tool Gives Housing To ({part})")
    doc.add_picture(str(f_), width=Inches(6.0))
    note("Characteristics are VI-SPDAT indicators written into each base profile. Housing = a referral to Permanent "
         "Supportive Housing (score 8 or more); truly needs housing = true score 8 or more. Significance tests are "
         "in Table 1.")
    doc.add_page_break()

heading("Table", 1, "Share of Each Group That Truly Needs Housing and That Each Tool Gives Housing To")
heads = ["Characteristic", "Group", "Records", "Truly needs housing", "VI-SPDAT gives housing", f"{AI} gives housing",
         "VI-SPDAT p", f"{AI} p"]
t = doc.add_table(rows=1, cols=len(heads))
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(heads):
    c = t.rows[0].cells[i]
    c.text = ""
    c.paragraphs[0].add_run(h).font.size = Pt(9)
    rule(c, "top", 8)
    rule(c, "bottom", 4)
prev = None
allrows = table.merge(tests, on=["section", "characteristic"])
for n, (_, r) in enumerate(allrows.iterrows()):
    first = r.characteristic != prev
    pv = r.p_vis_holm if r.section == "Profile" else r.p_vis
    pa = r.p_ai_holm if r.section == "Profile" else r.p_ai
    vals = [r.characteristic if first else "", r.group, str(r.records), f"{100 * r.truly:.1f}%",
            f"{100 * r['VI-SPDAT']:.1f}%", f"{100 * r[AI]:.1f}%",
            (("—" if r.section == "Demographic" and r.characteristic != "Disclosure" else fp(pv).replace("= ", "")) if first else ""),
            (fp(pa).replace("= ", "") if first else "")]
    prev = r.characteristic
    cells = t.add_row().cells
    for i, v in enumerate(vals):
        cells[i].text = ""
        run = cells[i].paragraphs[0].add_run(v)
        run.font.size = Pt(9)
        cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if i < 2 else WD_ALIGN_PARAGRAPH.CENTER
        if n == len(allrows) - 1:
            rule(cells[i], "bottom", 8)
for row in t.rows:
    for c, w in zip(row.cells, [1.6, 1.4, 0.55, 0.75, 0.75, 0.75, 0.6, 0.6]):
        c.width = Inches(w)
note("p tests whether the tool's share given housing departs from the share truly needing it differently across "
     "the groups of that characteristic (a broad test of over- and under-allocation combined; it does not "
     "separately test over-referral of people who are not truly high-need). Demographic characteristics: permutation within base profile; profile "
     "characteristics: permutation across base profiles, Holm-corrected. — : the VI-SPDAT does not use race, "
     "gender, or city, so its shares are identical across those groups.")
doc.add_page_break()
heading("Table", 2, "Share of All Housing Going to Each Group")
heads2 = ["Characteristic", "Group", "Share of population", "Share of everyone who truly needs housing",
          "Share of VI-SPDAT's housing", f"Share of {AI}'s housing", f"{AI} p"]
t2 = doc.add_table(rows=1, cols=len(heads2))
t2.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(heads2):
    c = t2.rows[0].cells[i]
    c.text = ""
    c.paragraphs[0].add_run(h).font.size = Pt(9)
    rule(c, "top", 8)
    rule(c, "bottom", 4)
prev = None
for n, (_, r) in enumerate(allrows.iterrows()):
    first = r.characteristic != prev
    pa = r.p_ai_holm if r.section == "Profile" else r.p_ai
    vals = [r.characteristic if first else "", r.group, f"{100 * r.pop_share:.1f}%", f"{100 * r.need_share:.1f}%",
            f"{100 * r.vis_share:.1f}%", f"{100 * r.ai_share:.1f}%", fp(pa).replace("= ", "") if first else ""]
    prev = r.characteristic
    cells = t2.add_row().cells
    for i, v in enumerate(vals):
        cells[i].text = ""
        cells[i].paragraphs[0].add_run(v).font.size = Pt(9)
        cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if i < 2 else WD_ALIGN_PARAGRAPH.CENTER
        if n == len(allrows) - 1:
            rule(cells[i], "bottom", 8)
for row in t2.rows:
    for c, w in zip(row.cells, [1.6, 1.4, 0.75, 0.9, 0.85, 0.85, 0.6]):
        c.width = Inches(w)
note("Each column shows the group's share of a whole: of all 640 records, of the 280 that truly need housing, and "
     f"of the records each tool gives housing to (VI-SPDAT 220; {AI} {int(TOOLS[AI].sum())}). Within each "
     "characteristic the groups add up to 100%. For race, gender, city, and disclosure every group is the same size "
     "and has the same real need, so a fair tool gives each group its population share. For profile "
     "characteristics, which differ in real need, the fair benchmark is the share of everyone who truly needs "
     "housing. p as in Table 1.")
out_path = OUT / (args.out or f"allocation_comparison_{args.run}.docx")
doc.save(out_path)
print(f"wrote {out_path.relative_to(ROOT)}")
for _, r in allrows.iterrows():
    print(f"{r.characteristic[:34]:34} {r.group[:32]:32} n={r.records:3}  truly {100*r.truly:5.1f}%  VI {100*r['VI-SPDAT']:5.1f}%  {AI} {100*r[AI]:5.1f}%")
print(tests[["characteristic", "p_vis_holm", "p_ai_holm"]].to_string())
