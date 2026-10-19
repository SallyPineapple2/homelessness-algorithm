"""
Export the base-profile table for use in a paper.

Scores are computed from each profile's indicator flags with the published VI-SPDAT scoring
rules, the same way the dashboard computes them, and checked against the VI-SPDAT results
file (full disclosure) before anything is written.

Writes: exports/base_profiles_table.csv   (opens in Excel, Google Sheets, or Word)
        exports/base_profiles_table.docx  (APA-style Word table, ready to paste)

Run:  python scripts/export_profile_table.py
"""

import ast
import csv
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "exports"

profiles = json.loads((ROOT / "data" / "base_profiles.json").read_text(encoding="utf-8"))["profiles"]
instrument = json.loads((ROOT / "data" / "vispdat_instrument.json").read_text(encoding="utf-8"))
domain_of = {ind["id"]: d["key"] for d in instrument["domains"] for ind in d["indicators"]}
name_of = {ind["id"]: ind["name"] for d in instrument["domains"] for ind in d["indicators"]}

# indicators withheld under underdisclosure, confirmed identical across each profile's 50 clones
withheld = {}
for c in json.loads((ROOT / "data" / "clones.json").read_text(encoding="utf-8"))["clones"]:
    iv = c["interview"] if isinstance(c["interview"], dict) else ast.literal_eval(c["interview"])
    w = tuple(iv["Underdisclosure"]["withheld"])
    assert withheld.setdefault(c["base_profile"], w) == w, c["base_profile"]
bands = instrument["bands"]

# cross-check against the scored VI-SPDAT results
scored = {}
with (ROOT / "data" / "vispdat_results.csv").open(encoding="utf-8-sig") as fh:
    for r in csv.DictReader(fh):
        if r["disclosure"] == "Full disclosure":
            scored[r["profile"]] = int(r["score"])

rows = []
for p in profiles:
    by = {k: sum(1 for i in p["indicators"] if domain_of.get(i) == k) for k in ("pre", "a", "b", "c", "d")}
    total = len(p["indicators"])
    assert total == sum(by.values()), p["id"]
    assert total == scored[p["id"]], f'{p["id"]}: {total} vs scored {scored[p["id"]]}'
    band = next(b for b in bands if b["min"] <= total <= b["max"])
    rows.append({
        "ID": p["id"], "Profile": p["label"], "Age": p["age"],
        # the age has its own column, so drop a leading "Age 68." from the summary
        "Summary": re.sub(r"^Age \d+\.\s*", "", p["vignette"]),
        "Pre-survey": by["pre"], "A. Housing history": by["a"], "B. Risks": by["b"],
        "C. Socialization and daily functions": by["c"], "D. Wellness": by["d"],
        "Total": total, "Recommended intervention": band["name"],
        "Withheld under underdisclosure": "; ".join(name_of[i] for i in withheld[p["id"]]) or "None",
    })

OUT.mkdir(exist_ok=True)
with (OUT / "base_profiles_table.csv").open("w", encoding="utf-8-sig", newline="") as fh:
    w = csv.DictWriter(fh, fieldnames=list(rows[0]))
    w.writeheader()
    w.writerows(rows)

# ---------- Word table, APA 7 style: horizontal rules only, no vertical lines ----------
doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Inches(0.75))
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(11)


def para(text, bold=False, italic=False, size=None, after=0):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.bold, run.italic = bold, italic
    if size:
        run.font.size = Pt(size)
    par.paragraph_format.space_after = Pt(after)
    return par


def rule(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:color"), "000000")
    borders.append(el)


para("Table 1", bold=True, after=0)
para("Base Profiles: Summary, VI-SPDAT Domain Scores, and Underdisclosure", italic=True, after=6)

heads = ["ID", "Profile", "Age", "Summary", "Pre", "A", "B", "C", "D", "Total", "Recommended intervention", "Withheld under underdisclosure"]
keys = ["ID", "Profile", "Age", "Summary", "Pre-survey", "A. Housing history", "B. Risks",
        "C. Socialization and daily functions", "D. Wellness", "Total", "Recommended intervention",
        "Withheld under underdisclosure"]
table = doc.add_table(rows=1, cols=len(heads))
table.alignment = WD_TABLE_ALIGNMENT.CENTER
widths = [0.4, 1.15, 0.38, 3.0, 0.32, 0.3, 0.3, 0.3, 0.3, 0.42, 1.05, 1.08]
numeric = {"Age", "Pre", "A", "B", "C", "D", "Total"}
for i, h in enumerate(heads):
    c = table.rows[0].cells[i]
    c.text = ""
    run = c.paragraphs[0].add_run(h)
    run.font.size = Pt(8.5)
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if h in numeric else WD_ALIGN_PARAGRAPH.LEFT
    rule(c, "top", 8)
    rule(c, "bottom", 4)
for n, r in enumerate(rows):
    cells = table.add_row().cells
    for i, (h, k) in enumerate(zip(heads, keys)):
        cells[i].text = ""
        run = cells[i].paragraphs[0].add_run(str(r[k]))
        run.font.size = Pt(8.5)
        cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if h in numeric else WD_ALIGN_PARAGRAPH.LEFT
        if n == len(rows) - 1:
            rule(cells[i], "bottom", 8)
for row in table.rows:
    for i, w in enumerate(widths):
        row.cells[i].width = Inches(w)
        row.cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(6)
note.add_run("Note. ").italic = True
note.add_run(
    "Scores are indicator counts under the VI-SPDAT for Single Adults, American Version 2.0 "
    "(OrgCode Consulting, 2015). Pre = pre-survey age indicator (maximum 1); A = History of "
    "Housing and Homelessness (maximum 2); B = Risks (maximum 4); C = Socialization and Daily "
    "Functions (maximum 4); D = Wellness (maximum 6). Total ranges from 0 to 17. Recommended "
    "intervention follows the published bands: 0–3, no housing intervention; 4–7, assessment for "
    "Rapid Re-Housing; 8–17, assessment for Permanent Supportive Housing / Housing First. Under "
    "underdisclosure, the listed indicators are withheld and the total falls by one point for each. "
    "Each profile was cloned across 2 genders, 5 races, and 5 cities (50 versions per profile); full "
    "narratives are in Appendix B."
).font.size = Pt(10)
doc.save(OUT / "base_profiles_table.docx")
print(f"wrote {len(rows)} profiles to exports/base_profiles_table.csv and .docx (totals match the scored results)")
