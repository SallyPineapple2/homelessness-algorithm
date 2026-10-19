"""
Export the paper appendix: how the demographic versions were built, and every base profile's
full narrative with its underdisclosure text.

Appendix A  Demographic versions: the 2 x 5 x 5 design, an example of what a model reads,
            and where the full list of 1,600 clones is published.
Appendix B  Base profile narratives: for each of the 32 profiles, its score, flagged
            indicators, the eleven narrative fields, and (if any) the indicators withheld under
            underdisclosure with the replacement text on record.

Everything is read from the same data files as the dashboard; the withheld indicators are
checked to be identical across all 50 clones of each profile before anything is written.

Writes: exports/appendix.docx
Run:    python scripts/export_appendix.py
"""

import ast
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "exports"
REPO = "https://github.com/SallyPineapple2/homelessness-algorithm"

profiles = json.loads((ROOT / "data" / "base_profiles.json").read_text(encoding="utf-8"))["profiles"]
clones_doc = json.loads((ROOT / "data" / "clones.json").read_text(encoding="utf-8"))
clones = clones_doc["clones"]
instrument = json.loads((ROOT / "data" / "vispdat_instrument.json").read_text(encoding="utf-8"))
ind = {i["id"]: (d, i) for d in instrument["domains"] for i in d["indicators"]}
bands = instrument["bands"]

FIELDS = [
    ("background", "Work and education"), ("path", "Path into homelessness"),
    ("housing", "Current housing"), ("finances", "Finances"), ("health", "Health"),
    ("safety", "Safety"), ("ties", "Family and social ties"), ("routine", "A typical day"),
    ("goals", "Goals"), ("services", "Service history"), ("demeanor", "Presentation in the interview"),
]
REGION = {"Los Angeles, California": "West", "Cincinnati, Ohio": "Midwest", "Chicago, Illinois": "Midwest",
          "Atlanta, Georgia": "South", "New York City, New York": "Northeast"}


def interview(clone):
    iv = clone["interview"]
    return iv if isinstance(iv, dict) else ast.literal_eval(iv)


# withheld indicators per profile, confirmed identical across its 50 clones
withheld = {}
for p in profiles:
    sets = {tuple(interview(c)["Underdisclosure"]["withheld"]) for c in clones if c["base_profile"] == p["id"]}
    assert len(sets) == 1, f'{p["id"]}: withheld indicators differ across clones'
    withheld[p["id"]] = list(sets.pop())
assert len(clones) == 50 * len(profiles)

doc = Document()
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(doc.sections[0], side, Inches(1))
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)
for name in ("Heading 1", "Heading 2", "Heading 3"):
    st = doc.styles[name]
    st.font.name = "Times New Roman"
    st.font.color.rgb = None
    st.font.size = Pt({"Heading 1": 14, "Heading 2": 12, "Heading 3": 11}[name])


def para(text="", bold=False, italic=False, size=None, after=6, before=0):
    par = doc.add_paragraph()
    if text:
        run = par.add_run(text)
        run.bold, run.italic = bold, italic
        if size:
            run.font.size = Pt(size)
    par.paragraph_format.space_after = Pt(after)
    par.paragraph_format.space_before = Pt(before)
    return par


def labeled(label, text, after=3):
    par = doc.add_paragraph()
    par.add_run(f"{label}. ").italic = True
    par.add_run(text)
    par.paragraph_format.space_after = Pt(after)
    return par


def rule(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:color"), "000000")
    borders.append(el)


def apa_table(heads, rows, widths):
    t = doc.add_table(rows=1, cols=len(heads))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(heads):
        c = t.rows[0].cells[i]
        c.text = ""
        c.paragraphs[0].add_run(h).font.size = Pt(10)
        rule(c, "top", 8)
        rule(c, "bottom", 4)
    for n, r in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            cells[i].paragraphs[0].add_run(str(v)).font.size = Pt(10)
            if n == len(rows) - 1:
                rule(cells[i], "bottom", 8)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    return t


# ---------------- Appendix A ----------------
doc.add_heading("Appendix A: Demographic Versions", level=1)
para(
    "Each of the 32 base profiles was cloned into 50 demographic versions by fully crossing two "
    "genders, the five race and ethnicity categories used by the U.S. Department of Housing and "
    "Urban Development, and five cities (2 × 5 × 5 = 50), for 1,600 clones. Clones of the same "
    "profile share an identical narrative; only the demographic line differs. Each clone was "
    "scored under full disclosure and under underdisclosure, giving 3,200 instances.")
para("Table A1", bold=True, after=0, before=6)
para("Factors Crossed to Create the Demographic Versions", italic=True, after=4)
genders = sorted({c["gender"] for c in clones}, key=["Male", "Female"].index)
races = list(dict.fromkeys(c["race"] for c in clones))
locs = clones_doc["locations"]
n = max(len(genders), len(races), len(locs))
rows = [[genders[i] if i < len(genders) else "", races[i] if i < len(races) else "",
         f"{locs[i]} ({REGION[locs[i]]})" if i < len(locs) else ""] for i in range(n)]
apa_table(["Gender (2)", "Race and ethnicity (5)", "City (5)"], rows, [1.3, 2.9, 2.3])
para("", after=4)

para("Example", italic=True, after=2, before=6)
ex = next(c for c in clones if c["base_profile"] == "P21" and c["gender"] == "Female"
          and c["race"].startswith("Black") and c["location"].startswith("Chicago"))
para("Every case a model read was labeled only by its position in the session (for example, "
     "“Case 7”) and opened with a demographic line in this form, followed by the "
     "profile's eleven narrative fields:")
q = para(ex["demographic_line"], after=6)
q.paragraph_format.left_indent = Inches(0.5)
para(
    f"The complete list of all 1,600 clones, including every demographic line and the record "
    f"under each disclosure condition, is available as supplementary data (data/clones.json) at "
    f"{REPO}.")

# ---------------- Appendix B ----------------
doc.add_page_break()
doc.add_heading("Appendix B: Base Profile Narratives", level=1)
para(
    "Each base profile is listed with its VI-SPDAT total under full disclosure, the recommended "
    "intervention for that total, and the vulnerability indicators it carries. The eleven "
    "narrative fields are the record every model read. Under underdisclosure, the listed "
    "indicators are withheld and the named fields are replaced with the text shown, which is what "
    "an assessor would have on record if the person never mentioned them. Scores follow the "
    "VI-SPDAT for Single Adults, American Version 2.0 (OrgCode Consulting, 2015).")

for p in profiles:
    total = len(p["indicators"])
    band = next(b for b in bands if b["min"] <= total <= b["max"])
    doc.add_heading(f'{p["id"]}. {p["label"]}', level=2)
    para(f'Age {p["age"]} · VI-SPDAT total {total} of 17 · {band["name"]}', italic=True, after=4)
    flagged = [f'{ind[i][1]["name"]} ({ind[i][0]["name"]})' for i in p["indicators"]]
    labeled("Indicators flagged", "; ".join(flagged) + "." if flagged else "None.", after=6)
    for key, label in FIELDS:
        labeled(label, p["narrative"][key])
    w = withheld[p["id"]]
    under = p.get("underdisclosure") or {}
    if w:
        par = para("", after=3, before=6)
        par.add_run("Underdisclosure").bold = True
        names = [f'{ind[i][1]["name"]}' for i in w]
        labeled("Indicators withheld", "; ".join(names) + f" (total falls to {total - len(w)}).")
        for key, label in FIELDS:
            if key in under:
                labeled(f"{label}, as recorded", under[key])
    else:
        labeled("Underdisclosure", "No indicators withheld; the record is the same under both conditions.", after=6)

OUT.mkdir(exist_ok=True)
doc.save(OUT / "appendix.docx")
print(f"wrote exports/appendix.docx: {len(profiles)} profiles, {len(clones)} clones summarized")
