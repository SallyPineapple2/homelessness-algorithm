"""
Where do the AI runs score higher or lower than the VI-SPDAT, and does the pattern follow gender,
race, city, or other characteristics of the person?

For every scored record, gap = AI score - VI-SPDAT score on the same record (same profile,
same demographic version, same disclosure condition).

  1. Direction and size: how often each AI run scores lower, the same, or higher, and by how much.
  2. Patterns by characteristic: one regression per AI run, gap on race, gender, city, disclosure,
     and true vulnerability band, with standard errors clustered by base profile
     (generalized estimating equations, exchangeable correlation), since each profile appears 20
     times. Joint Wald tests for race (4 df) and city (4 df).
  3. Which vulnerability indicators widen or narrow the gap (exploratory, profile level, full
     disclosure): per-profile mean gap regressed on the profile's true total and the indicator.
  4. Figure: mean gap at each true score, by disclosure condition.

Writes: data/ai/results/gap_patterns_*.csv, exports/gap_patterns.docx,
        exports/figures/figure4_gap_by_true_score.png
Run:    python scripts/analyze_gap_patterns.py
"""

import csv
import json
import warnings
from collections import defaultdict
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import statsmodels.api as sm
import statsmodels.formula.api as smf
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

warnings.filterwarnings("ignore")
ROOT = Path(__file__).resolve().parents[1]
RES = ROOT / "data" / "ai" / "results"
OUT = ROOT / "exports"
(OUT / "figures").mkdir(parents=True, exist_ok=True)

roster = json.loads((ROOT / "data" / "ai" / "models.json").read_text(encoding="utf-8"))["models"]
import argparse
_ap = argparse.ArgumentParser()
_ap.add_argument("--runs", help="comma-separated AI run keys, e.g. claude (default: every run)")
_ap.add_argument("--out", default="gap_analysis_for_professor.docx", help="file name for the shareable document")
_ap.add_argument("--data-only", action="store_true",
                 help="tables and figure only: no summary or method text, and notes limited to reading the numbers")
ARGS = _ap.parse_args()
RUNS = [(m["label"], m["key"]) for m in roster if (RES / f'{m["key"]}.csv').exists()
        and (not ARGS.runs or m["key"] in ARGS.runs.split(","))]
SINGLE = len(RUNS) == 1
profiles = {p["id"]: p for p in json.loads((ROOT / "data" / "base_profiles.json").read_text(encoding="utf-8"))["profiles"]}
inst = json.loads((ROOT / "data" / "vispdat_instrument.json").read_text(encoding="utf-8"))
IND = [(i["id"], i["name"], d["name"]) for d in inst["domains"] for i in d["indicators"]]

SHORT = {"American Indian, Alaska Native, or Indigenous": "AIAN", "Asian or Asian American": "Asian",
         "Black, African American, or African": "Black", "Hispanic/Latino/e/a": "Hispanic", "White": "White"}
CITY = {"Los Angeles, California": "LA", "Cincinnati, Ohio": "Cincinnati", "Chicago, Illinois": "Chicago",
        "Atlanta, Georgia": "Atlanta", "New York City, New York": "NYC"}


def band(t):
    return "0-3" if t <= 3 else "4-7" if t <= 7 else "8-12" if t <= 12 else "13-17"


def load(key):
    df = pd.DataFrame(list(csv.DictReader((RES / f"{key}.csv").open(encoding="utf-8-sig"))))
    for c in ("score", "vispdat_score", "true_score"):
        df[c] = df[c].astype(int)
    df["gap"] = df["score"] - df["vispdat_score"]
    df["race_s"] = df["race"].map(SHORT)
    df["city_s"] = df["location"].map(CITY)
    df["under"] = (df["disclosure"] == "Underdisclosure").astype(int)
    df["female"] = (df["gender"] == "Female").astype(int)
    df["tband"] = df["true_score"].map(band)
    return df


def holm(ps):
    order = np.argsort(ps)
    adj, run = [0.0] * len(ps), 0.0
    for rank, i in enumerate(order):
        run = max(run, min(1.0, (len(ps) - rank) * ps[i]))
        adj[i] = run
    return adj


def fp(p):
    if p < 0.001:
        return "< .001"
    s3 = f"{p:.3f}".lstrip("0")
    # when rounding to 3 places hides which side of .05 the value is on, show 4
    if s3 != ".050":
        return s3
    s4 = f"{p:.4f}".lstrip("0")
    return f"{p:.5f}".lstrip("0") if s4 == ".0500" else s4


data = {label: load(key) for label, key in RUNS}

# ---------- 1. direction and size ----------
CATS = [("AI lower", lambda g: g < 0), ("Same", lambda g: g == 0), ("Higher by 1–2", lambda g: (g >= 1) & (g <= 2)),
        ("Higher by 3–5", lambda g: (g >= 3) & (g <= 5)), ("Higher by 6+", lambda g: g >= 6)]
dist_rows = []
for label, df in data.items():
    for sub_name, sub in (("All records", df), ("Full disclosure", df[df.under == 0]), ("Underdisclosure", df[df.under == 1])):
        row = {"run": label, "records": sub_name, "n": len(sub)}
        for name, f in CATS:
            row[name] = f(sub.gap).mean()
        row["mean_gap"] = sub.gap.mean()
        dist_rows.append(row)
if not ARGS.runs:
    pd.DataFrame(dist_rows).round(4).to_csv(RES / "gap_patterns_distribution.csv", index=False, encoding="utf-8-sig")

# ---------- 2. clustered regression ----------
FORMULA = ("gap ~ C(race_s, Treatment('White')) + female + C(city_s, Treatment('Atlanta')) + under"
           " + C(tband, Treatment('0-3'))")
reg = {}
for label, df in data.items():
    model = smf.gee(FORMULA, groups="profile", data=df, cov_struct=sm.cov_struct.Exchangeable(),
                    family=sm.families.Gaussian())
    # Mancl-DeRouen bias-reduced robust covariance: only 32 clusters
    res = model.fit(cov_type="bias_reduced")
    names = list(res.params.index)
    joint = {}
    for fam, tag in (("race", "race_s"), ("city", "city_s")):
        idx = [i for i, n in enumerate(names) if tag in n]
        L = np.zeros((len(idx), len(names)))
        for r, i in enumerate(idx):
            L[r, i] = 1
        w = res.wald_test(L, scalar=True)
        joint[fam] = (float(w.statistic), len(idx), float(w.pvalue))
    reg[label] = (res, joint)

TERMS = [
    ("Race (vs. White)", None), ("   American Indian or Alaska Native", "C(race_s, Treatment('White'))[T.AIAN]"),
    ("   Asian", "C(race_s, Treatment('White'))[T.Asian]"), ("   Black", "C(race_s, Treatment('White'))[T.Black]"),
    ("   Hispanic or Latino", "C(race_s, Treatment('White'))[T.Hispanic]"), ("   Joint test, χ²(4)", "joint:race"),
    ("Female (vs. male)", "female"),
    ("City (vs. Atlanta)", None), ("   Chicago", "C(city_s, Treatment('Atlanta'))[T.Chicago]"),
    ("   Cincinnati", "C(city_s, Treatment('Atlanta'))[T.Cincinnati]"), ("   Los Angeles", "C(city_s, Treatment('Atlanta'))[T.LA]"),
    ("   New York City", "C(city_s, Treatment('Atlanta'))[T.NYC]"), ("   Joint test, χ²(4)", "joint:city"),
    ("Underdisclosure (vs. full)", "under"),
    ("True vulnerability (vs. 0–3)", None), ("   4–7", "C(tband, Treatment('0-3'))[T.4-7]"),
    ("   8–12", "C(tband, Treatment('0-3'))[T.8-12]"), ("   13–17", "C(tband, Treatment('0-3'))[T.13-17]"),
    ("Intercept", "Intercept"),
]


def reg_cell(label, term):
    res, joint = reg[label]
    if term is None:
        return ""
    if term.startswith("joint:"):
        stat, df_, p = joint[term.split(":")[1]]
        return f"{stat:.2f}, p {fp(p).replace('< ', '<')}" if p < 0.001 else f"{stat:.2f}, p = {fp(p)}"
    b, (lo, hi), p = res.params[term], res.conf_int().loc[term], res.pvalues[term]
    star = "*" if p < 0.05 else ""
    return f"{b:+.2f}{star} [{lo:+.2f}, {hi:+.2f}]"


reg_csv = []
for label in data:
    res, joint = reg[label]
    for term in res.params.index:
        lo, hi = res.conf_int().loc[term]
        reg_csv.append({"run": label, "term": term, "coef": res.params[term], "ci_low": lo, "ci_high": hi, "p": res.pvalues[term]})
    for fam, (stat, df_, p) in joint.items():
        reg_csv.append({"run": label, "term": f"joint {fam} ({df_} df)", "coef": stat, "ci_low": "", "ci_high": "", "p": p})
if not ARGS.runs:
    pd.DataFrame(reg_csv).to_csv(RES / "gap_patterns_regression.csv", index=False, encoding="utf-8-sig")

# ---------- 3. indicators that widen or narrow the gap (profile level, full disclosure) ----------
ind_rows = []
prof_gap = {}
for label, df in data.items():
    full = df[df.under == 0]
    pg = full.groupby("profile").gap.mean()
    prof_gap[label] = pg
pmean = pd.concat(prof_gap, axis=1).mean(axis=1)  # average across AI runs
base = pd.DataFrame({"gap": pmean, "true": [len(profiles[p]["indicators"]) for p in pmean.index]}, index=pmean.index)
pvals, coefs = [], []
# in this profile set, Mental Health and Tri-Morbidity flag exactly the same 8 profiles, so they are one test
same = {p for p in profiles if "d3_mental" in profiles[p]["indicators"]} == {p for p in profiles if "d4_trimorbidity" in profiles[p]["indicators"]}
for iid, iname, dname in IND:
    if same and iid == "d4_trimorbidity":
        continue
    if same and iid == "d3_mental":
        iname = "Mental Health (always with Tri-Morbidity)"
    base["has"] = [int(iid in profiles[p]["indicators"]) for p in base.index]
    n_with = int(base["has"].sum())
    if n_with in (0, len(base)):
        continue
    fit = smf.ols("gap ~ true + has", data=base).fit()
    coefs.append((iid, iname, dname, n_with, fit.params["has"], *fit.conf_int().loc["has"], fit.pvalues["has"],
                  base.loc[base.has == 1, "gap"].mean(), base.loc[base.has == 0, "gap"].mean()))
    pvals.append(fit.pvalues["has"])
for row, pa in zip(coefs, holm(pvals)):
    ind_rows.append({"indicator": row[1], "domain": row[2], "profiles_with": row[3], "coef_adj_true": row[4],
                     "ci_low": row[5], "ci_high": row[6], "p": row[7], "p_holm": pa,
                     "mean_gap_with": row[8], "mean_gap_without": row[9]})
ind_df = pd.DataFrame(ind_rows).sort_values("coef_adj_true", ascending=False)
if not ARGS.runs:
    ind_df.round(4).to_csv(RES / "gap_patterns_indicators.csv", index=False, encoding="utf-8-sig")

# profiles with the largest and smallest gaps (average across AI runs, full disclosure)
ext = pd.DataFrame({"profile": pmean.index, "label": [profiles[p]["label"] for p in pmean.index],
                    "true": base["true"].values, "gap": pmean.values}).sort_values("gap")

# ---------- 4. figure ----------
plt.rcParams.update({"font.family": "Times New Roman", "font.size": 10, "axes.spines.top": False,
                     "axes.spines.right": False, "savefig.dpi": 300, "savefig.bbox": "tight"})
markers = ["o", "s", "^", "D", "v"]
fig, axes = plt.subplots(1, 2, figsize=(7.5, 3.2), sharey=True)
for ax, (cond, title) in zip(axes, ((0, "Full disclosure"), (1, "Underdisclosure"))):
    for (label, df), mk in zip(data.items(), markers):
        sub = df[df.under == cond].groupby("true_score").gap.mean()
        ax.plot(sub.index, sub.values, marker=mk, markersize=4, linewidth=1, color="black",
                markerfacecolor="white" if mk in "s^" else "black", label=label)
    ax.axhline(0, color="black", linestyle=":", linewidth=1)
    ax.set_title(title, loc="left", fontsize=10)
    ax.set_xlabel("True vulnerability score")
    ax.set_xticks(range(0, 18, 2))
axes[0].set_ylabel(f"{RUNS[0][0] if SINGLE else 'AI'} minus VI-SPDAT (points)")
if not SINGLE:
    fig.legend(*axes[0].get_legend_handles_labels(), frameon=False, fontsize=9, loc="lower center", ncol=4, bbox_to_anchor=(0.5, -0.06))
for ax, letter in zip(axes, "AB"):
    ax.text(-0.02, 1.04, letter, transform=ax.transAxes, fontsize=11, fontweight="bold", ha="right")
fig.tight_layout()
FIG_PATH = OUT / "figures" / ("figure4_gap_by_true_score.png" if not ARGS.runs else f"gap_by_true_score_{ARGS.runs.replace(',', '_')}.png")
fig.savefig(FIG_PATH)
plt.close(fig)

# ---------- Word files ----------
doc = None

def rule(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:color"), "000000")
    b.append(el)


def heading(kind, number, title):
    p = doc.add_paragraph()
    p.add_run(f"{kind} {number}").bold = True
    p.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph()
    p.add_run(title).italic = True
    p.paragraph_format.space_after = Pt(6)


DATA_ONLY = False
WHO = "AI"


def note(text, short=None):
    if DATA_ONLY and short is not None:
        text = short
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.add_run("Note. ").italic = True
    p.add_run(text).font.size = Pt(10)
    doc.add_paragraph()


def table(heads, rows, widths, size=9.5):
    t = doc.add_table(rows=1, cols=len(heads))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(heads):
        c = t.rows[0].cells[i]
        c.text = ""
        c.paragraphs[0].add_run(h).font.size = Pt(size)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
        rule(c, "top", 8)
        rule(c, "bottom", 4)
    for n, r in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(size)
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            if n == len(rows) - 1:
                rule(cells[i], "bottom", 8)
    for row in t.rows:
        for c, w in zip(row.cells, widths):
            c.width = Inches(w)


def new_doc():
    global doc
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.9))
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(12)


def add_items(order):
    """order: list of (kind, number) for the distribution table, the figure, the regression, and the indicators."""
    for item, number in order:
        if item == "dist":
            heading("Table", number, f"How Often {WHO if SINGLE else 'Each AI Run'} Scored Lower Than, Equal to, or Higher Than the VI-SPDAT")
            rows = []
            prev = None
            for r in dist_rows:
                rows.append([r["run"] if r["run"] != prev else "", r["records"], r["n"]] +
                            [f'{100 * r[c]:.1f}%' for c, _ in CATS] + [f'{r["mean_gap"]:+.2f}'])
                prev = r["run"]
            table(["AI run", "Records", "n"] + [c for c, _ in CATS] + ["Mean gap"], rows, [1.6, 1.3, 0.45, 0.8, 0.7, 1.0, 1.0, 0.95, 0.8])
            note("Gap is the AI score minus the VI-SPDAT score on the same record, in points on the 0–17 scale. "
                 "Each AI run scored 640 records (320 under each disclosure condition).",
                 short=f"Gap = {WHO} score minus VI-SPDAT score on the same record, in points (0–17 scale). "
                       "n = number of records.")

        elif item == "regression":
            labels = list(data)
            heading("Table", number, f"Characteristics Associated With the Size of the {WHO}–VI-SPDAT Gap")
            rows = [[name] + [reg_cell(l, term) for l in labels] for name, term in TERMS]
            table(["Characteristic"] + labels, rows, [2.4] + [1.75] * len(labels))
            note("Coefficients are the change in the gap (AI minus VI-SPDAT, points) associated with each "
                 "characteristic, holding the others constant, with 95% confidence intervals; * p < .05. Linear "
                 "models estimated by generalized estimating equations with an exchangeable working correlation and "
                 "bias-reduced (Mancl–DeRouen) robust standard errors clustered by base profile (32 clusters, 20 "
                 "records each). Joint tests are "
                 "Wald χ² tests that all race (or all city) coefficients are zero. True vulnerability is a "
                 "property of the base profile; race, gender, city, and disclosure vary within it."
                 + ("" if SINGLE else " Across the four AI runs, only Gemini's joint race test remains significant "
                    "after Holm correction (4 tests)."),
                 short="Values are coefficients in points with 95% confidence intervals; * p < .05. Reference "
                       "groups: White, male, Atlanta, full disclosure, true score 0–3. Joint "
                       "tests are Wald χ²(4).")

        elif item == "indicators":
            heading("Table", number, f"Vulnerability Indicators Associated With Larger or Smaller {WHO}–VI-SPDAT Gaps")
            rows = [[r["indicator"], r["domain"], r["profiles_with"], f'{r["mean_gap_with"]:.2f}', f'{r["mean_gap_without"]:.2f}',
                     f'{r["coef_adj_true"]:+.2f} [{r["ci_low"]:+.2f}, {r["ci_high"]:+.2f}]', fp(r["p"]), fp(r["p_holm"])]
                    for _, r in ind_df.iterrows()]
            table(["Indicator", "Domain", "Profiles with it", "Mean gap, with", "Mean gap, without",
                   "Difference adjusted for true score [95% CI]", "p", "Holm p"], rows, [1.9, 2.1, 0.8, 0.8, 0.8, 1.9, 0.55, 0.65])
            note("Exploratory, at the level of the 32 base profiles under full disclosure. The gap is averaged across "
                 + ("" if SINGLE else "the AI runs and ") + "each profile's demographic versions. The adjusted difference is the indicator's "
                 "coefficient in a regression of the profile's mean gap on its true VI-SPDAT total and the indicator, so "
                 "it shows how much more (positive) or less (negative) the AI scores profiles with that indicator than "
                 "the VI-SPDAT does, beyond what their total predicts. Holm correction across the indicators tested. "
                 "In this profile set Mental Health and Tri-Morbidity occur on the same 8 profiles (totals 9–17) and "
                 "cannot be separated; because those profiles sit near the top of the scale, where the AI cannot score "
                 "much higher, a negative value for them may reflect a ceiling rather than the indicator itself.",
                 short="Mean gap per base profile under full disclosure. The adjusted difference controls for the "
                       "profile's VI-SPDAT total. Holm p corrects across the indicators shown. Mental Health and "
                       "Tri-Morbidity occur on the same 8 profiles.")

        elif item == "pvalues":
            heading("Table", number, "Significance of Each Comparison")
            res, joint = reg[list(data)[0]]
            names_ = {"AIAN": "American Indian or Alaska Native", "Asian": "Asian", "Black": "Black",
                      "Hispanic": "Hispanic or Latino"}
            cities_ = {"NYC": "New York City", "LA": "Los Angeles", "Chicago": "Chicago", "Cincinnati": "Cincinnati"}
            rows = [["Race, all five groups (joint)", "—", f"χ²(4) = {joint['race'][0]:.2f}", fp(joint["race"][2]),
                     "Yes" if joint["race"][2] < 0.05 else "No"]]
            for code, nm in names_.items():
                t_ = "C(race_s, Treatment('White'))[T." + code + "]"
                rows.append([f"   {nm} vs. White", f"{res.params[t_]:+.2f}", f"z = {res.tvalues[t_]:.2f}",
                             fp(res.pvalues[t_]), "Yes" if res.pvalues[t_] < 0.05 else "No"])
            rows.append(["Female vs. male", f"{res.params['female']:+.2f}", f"z = {res.tvalues['female']:.2f}",
                         fp(res.pvalues["female"]), "Yes" if res.pvalues["female"] < 0.05 else "No"])
            rows.append(["City, all five cities (joint)", "—", f"χ²(4) = {joint['city'][0]:.2f}", fp(joint["city"][2]),
                         "Yes" if joint["city"][2] < 0.05 else "No"])
            for code, nm in cities_.items():
                t_ = "C(city_s, Treatment('Atlanta'))[T." + code + "]"
                rows.append([f"   {nm} vs. Atlanta", f"{res.params[t_]:+.2f}", f"z = {res.tvalues[t_]:.2f}",
                             fp(res.pvalues[t_]), "Yes" if res.pvalues[t_] < 0.05 else "No"])
            for term, nm in (("under", "Underdisclosure vs. full disclosure"),):
                rows.append([nm, f"{res.params[term]:+.2f}", f"z = {res.tvalues[term]:.2f}", fp(res.pvalues[term]),
                             "Yes" if res.pvalues[term] < 0.05 else "No"])
            table(["Comparison", "Difference (points)", "Test statistic", "p", "p < .05"], rows,
                  [3.2, 1.4, 1.4, 0.8, 0.8])
            note("", short="Differences and tests from the model in Table 2. Joint tests are Wald χ²(4); "
                           "single comparisons are Wald z tests.")

        elif item == "figure":
            heading("Figure", number, f"Gap Between {WHO} and VI-SPDAT Scores at Each Level of True Vulnerability")
            doc.add_picture(str(FIG_PATH), width=Inches(8.5))
            note(("Points" if SINGLE else "Each line is one AI run; points") + " are mean gaps (AI minus VI-SPDAT) across all records with that "
                 "true score. The dotted line marks no difference. (A) Full disclosure, where the VI-SPDAT equals the "
                 "true score. (B) Underdisclosure, where the VI-SPDAT loses the withheld indicators and the gap widens.",
                 short="Mean gap at each true score. Dotted line = no difference. (A) Full disclosure. "
                       "(B) Underdisclosure.")


# working file, numbered to follow Tables 1-3 and Figures 1-3 of the paper exports
new_doc()
add_items([("dist", 4), ("regression", 5), ("indicators", 6), ("figure", 4)])
try:
    if ARGS.runs:
        raise PermissionError  # one-run versions never replace the full working file
    doc.save(OUT / "gap_patterns.docx")
except PermissionError:
    print("skipped exports/gap_patterns.docx: it is open in Word; close it and rerun to update it")

# ---------- console summary ----------
labels = list(data)
print("DISTRIBUTION")
for r in dist_rows:
    print(f'  {r["run"]:24} {r["records"]:16} lower {100*r["AI lower"]:4.1f}%  same {100*r["Same"]:4.1f}%  +1-2 {100*r["Higher by 1–2"]:4.1f}%  +3-5 {100*r["Higher by 3–5"]:4.1f}%  +6 {100*r["Higher by 6+"]:4.1f}%  mean {r["mean_gap"]:+.2f}')
print("REGRESSION")
for name, term in TERMS:
    print(f"  {name:38} " + " | ".join(f"{reg_cell(l, term):26}" for l in labels))
print("INDICATORS (adjusted for true score)")
for _, r in ind_df.iterrows():
    print(f'  {r["indicator"]:28} n={r["profiles_with"]:2}  {r["coef_adj_true"]:+.2f} [{r["ci_low"]:+.2f},{r["ci_high"]:+.2f}] p={r["p"]:.3f} holm={r["p_holm"]:.3f}')
print("PROFILES, smallest and largest mean gap (all runs, full disclosure)")
for _, r in pd.concat([ext.head(5), ext.tail(5)]).iterrows():
    print(f'  {r["profile"]} {r["label"][:40]:40} true {r["true"]:2}  gap {r["gap"]:+.2f}')
print("wrote exports/gap_patterns.docx and exports/figures/figure4_gap_by_true_score.png")


# ---------- one shareable file: summary, then Table 1, Figure 1, Table 2, Table 3 ----------
def para(text, bold=False, italic=False, size=12, after=6, before=0, align=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold, run.italic = bold, italic
    run.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    if align:
        p.alignment = align
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text).font.size = Pt(11)
    p.paragraph_format.space_after = Pt(3)


dist_all = {r["run"]: r for r in dist_rows if r["records"] == "All records"}
higher = [1 - dist_all[l]["AI lower"] - dist_all[l]["Same"] for l in data]
lower = [dist_all[l]["AI lower"] for l in data]
same = [dist_all[l]["Same"] for l in data]
under = [reg[l][0].params["under"] for l in data]
top = ext.tail(3).iloc[::-1]
top5 = ext.tail(5)
all_physical = all("d1_physical" in profiles[p]["indicators"] for p in top5.profile)
sig_ind = ind_df[(ind_df.p_holm < 0.05) & (ind_df.coef_adj_true > 0)]


def rng(vals, f):
    lo, hi = min(vals), max(vals)
    return f(lo) if f(lo) == f(hi) else f"{f(lo)}–{f(hi)}"


def pct(v):
    return f"{100 * v:.0f}%"


def pct1(v):
    return f"{100 * v:.1f}%"


RACE_NAME = {"AIAN": "American Indian or Alaska Native", "Asian": "Asian", "Black": "Black", "Hispanic": "Hispanic or Latino"}
CITY_NAME = {"Chicago": "Chicago", "Cincinnati": "Cincinnati", "LA": "Los Angeles", "NYC": "New York City"}


def demographic_sentence(label):
    res, joint = reg[label]
    race_bits = []
    for code, name in RACE_NAME.items():
        term = "C(race_s, Treatment('White'))[T." + code + "]"
        if res.pvalues[term] < 0.05:
            race_bits.append(f"{name} cases {res.params[term]:+.2f} points relative to White cases".replace("-0", "−0"))
    rstat, _, rp = joint["race"]
    if rp < 0.05:
        race = f"Race: the joint test is significant (χ²(4) = {rstat:.2f}, p = {fp(rp)})"
        race += (", driven by " + "; ".join(race_bits) + ".") if race_bits else "."
    else:
        race = f"Race: no significant pattern (joint test p = {fp(rp)})."
    fpv = res.pvalues["female"]
    if fpv < 0.05:
        gender = f"Gender: women {res.params['female']:+.2f} points relative to men (p = {fp(fpv)})."
    else:
        gender = f"Gender: no significant difference ({res.params['female']:+.2f} points for women, p = {fp(fpv)})."
    city_bits = []
    for code, name in CITY_NAME.items():
        term = "C(city_s, Treatment('Atlanta'))[T." + code + "]"
        if res.pvalues[term] < 0.05:
            city_bits.append(f"{name} {res.params[term]:+.2f}".replace("-", "−"))
    cp = joint["city"][2]
    city = f"City: joint test p = {fp(cp)}" + (" (" + "; ".join(city_bits) + " relative to Atlanta)" if city_bits else "") + "."
    return " ".join([race, gender, city])


model_name = list(data)[0] if SINGLE else None
if ARGS.data_only:
    DATA_ONLY = True
    WHO = model_name or "AI"
    new_doc()
    title = f"{model_name} and VI-SPDAT Scores: Results" if SINGLE else "AI and VI-SPDAT Scores: Results"
    para(title, bold=True, size=16, after=14)
    order = [("dist", 1), ("figure", 1), ("regression", 2), ("indicators", 3)]
    if SINGLE:
        order.append(("pvalues", 4))
    add_items(order)
    doc.save(OUT / ARGS.out)
    print(f"wrote exports/{ARGS.out}")
    raise SystemExit
new_doc()
if SINGLE:
    para(f"Where {model_name} and VI-SPDAT Scores Differ: A Pilot Analysis", bold=True, size=16, after=2)
    para(f"Comparing {model_name}'s scores with the VI-SPDAT's on the same synthetic profiles, by gender, race, and other characteristics",
         italic=True, size=12, after=14)
else:
    para("Where AI and VI-SPDAT Scores Differ", bold=True, size=16, after=2)
    para("Comparing AI and VI-SPDAT scores on the same synthetic profiles, by gender, race, and other characteristics",
         italic=True, size=12, after=14)
para("Question", bold=True, after=2)
para("For the same synthetic records, when does the AI score higher or lower than the VI-SPDAT, and do those "
     "differences follow gender, race, city, or other characteristics of the person?", after=10)
para("Data and method", bold=True, after=2)
if SINGLE and model_name == "Gemini":
    who = ("Gemini (gemini-3.8-flash) was run through Google's Gemini API, which receives only the prompt (no app "
           "instructions, memory, or account context), with each session in a new chat at default settings. Gemini")
elif SINGLE:
    who = model_name
else:
    who = f"{['Zero', 'One', 'Two', 'Three', 'Four', 'Five', 'Six'][len(data)]} AI runs (" + "; ".join(data) + ")"
para(f"{who} and the published VI-SPDAT each scored the same 640 records: 32 base profiles, each presented under "
     "20 combinations of race, gender, and disclosure, with five cities balanced across them. The AI was never "
     "shown the VI-SPDAT's questions or scoring rules. For every record, the gap is the AI score minus the "
     "VI-SPDAT score (0–17 scale). Table 2 models the gap on all characteristics at once, with standard errors "
     "clustered by profile.", after=10)
para("Main results", bold=True, after=2)
subj = model_name if SINGLE else "The AI"
bullet(f"{subj} scored higher than the VI-SPDAT on {rng(higher, pct1)} of records, the same on {rng(same, pct1)}, "
       f"and lower on {rng(lower, pct1)} (Table 1).")
bullet("The gap is largest for people in the middle of the scale, near the threshold for Permanent Supportive "
       "Housing, and smallest at the top, where scores reach the maximum (Figure 1).")
bullet(f"Underdisclosure widens the gap by {rng(under, lambda v: f'{v:.1f}')} points (Table 2): the VI-SPDAT loses "
       "the withheld points, while the AI still reads the need in the narrative.")
if len(sig_ind):
    r0 = sig_ind.iloc[0]
    text = (f"{r0['indicator']} widens the gap the most: {r0['coef_adj_true']:+.1f} points beyond what the profile's "
            f"total predicts (Holm p {fp(r0['p_holm'])}; Table 3).")
    if r0["indicator"] == "Physical Health":
        text += " The VI-SPDAT counts any health problem as one point; the AI appears to weigh its severity."
    text += (" All five profiles with the largest gaps have a physical health condition, led by " if all_physical
             else " The largest gaps were ")
    text += "; ".join(f"{r['label'].lower() if all_physical else r['label']} ({r['gap']:+.1f})" for _, r in top.iterrows()) + "."
    bullet(text)
else:
    bullet("No single vulnerability indicator is associated with the gap after Holm correction (Table 3).")
if SINGLE:
    bullet(demographic_sentence(model_name))
else:
    for l in data:
        bullet(f"{l}. " + demographic_sentence(l))
para("", after=4)
doc.add_page_break()
add_items([("dist", 1), ("figure", 1), ("regression", 2), ("indicators", 3)])
doc.save(OUT / ARGS.out)
print(f"wrote exports/{ARGS.out}")
