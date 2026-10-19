"""
A short review sheet for the 100 pilot profiles: one row per person, sorted by
VI-SPDAT score, so plausibility, variety, and scoring can be checked at a glance.
The full text stays in profiles_for_review.docx for any row that looks off.

Run:  python pilot/scripts/export_review_sheet.py
Writes: pilot/exports/profiles_review_sheet.docx
"""

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from analyze import apa_table, para, style_doc  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
PILOT = ROOT / "pilot"

SHORT = {
    "pre_age": "60+", "a1_sleep": "unsheltered", "a2_duration": "long/repeat homeless",
    "b1_emergency": "4+ ER/police", "b2_harm": "harm", "b3_legal": "legal",
    "b4_exploitation": "exploited", "c1_money": "no income/debt", "c2_activity": "no activity",
    "c3_selfcare": "self-care", "c4_relationships": "relationship cause", "d1_physical": "physical",
    "d2_substance": "substance", "d3_mental": "mental", "d4_trimorbidity": "tri-morbid",
    "d5_medications": "meds", "d6_trauma": "trauma cause",
}
ORDER = list(SHORT)
CITY = {"Los Angeles, California": "LA", "Cincinnati, Ohio": "Cin", "Chicago, Illinois": "Chi",
        "Atlanta, Georgia": "Atl", "New York City, New York": "NYC"}
RACE = {"American Indian, Alaska Native, or Indigenous": "AIAN", "Asian or Asian American": "Asian",
        "Black, African American, or African": "Black", "Hispanic/Latino/e/a": "Hisp", "White": "White"}


def main():
    profiles = json.loads((PILOT / "data" / "profiles.json").read_text(encoding="utf-8"))["profiles"]
    rows = sorted(profiles, key=lambda p: (len(p["indicators"]), int(p["id"][1:])))
    doc = Document()
    style_doc(doc, landscape=True)
    para(doc, "Pilot Profiles: Review Sheet", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    para(doc, "One row per person, lowest VI-SPDAT score first. Check three things: (1) no two people read as "
              "the same story; (2) each label and age are believable together; (3) the items listed match "
              "what you would expect for that person. Mark any row that looks wrong and read its full text "
              "in profiles_for_review.docx. P01-P32 come from the original study.", size=10, space_after=8)
    body = []
    for p in rows:
        s = len(p["indicators"])
        pri = "High" if s >= 8 else "Medium" if s >= 4 else "Low"
        items = ", ".join(SHORT[i] for i in ORDER if i in p["indicators"]) or "none"
        body.append([p["id"], p["label"], p["age"], "F" if p["gender"] == "Female" else "M",
                     RACE[p["race"]], CITY[p["location"]], s, pri, items, ""])
    apa_table(doc, 1, "All 100 Profiles, Sorted by VI-SPDAT Score",
              ["ID", "Who", "Age", "Sex", "Race", "City", "Score", "Priority", "VI-SPDAT items met", "OK?"],
              body, size=8, widths=[0.4, 2.3, 0.35, 0.3, 0.45, 0.35, 0.4, 0.55, 3.6, 0.4],
              note="Items: unsheltered = not in a shelter; long/repeat homeless = 1+ year or 4+ episodes; "
                   "4+ ER/police = 4+ emergency contacts in 6 months; harm = attacked or threatened harm; "
                   "exploited = forced, tricked, or trading sex/drugs/favors; self-care = cannot meet basic "
                   "needs; relationship/trauma cause = homelessness caused by it; physical/mental/substance = "
                   "condition affecting housing; meds = not taking or selling prescriptions.")
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for para_ in c.paragraphs:
                    para_.paragraph_format.space_after = Pt(0)
    out = PILOT / "exports" / "profiles_review_sheet.docx"
    doc.save(out)
    print(f"wrote {out.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
