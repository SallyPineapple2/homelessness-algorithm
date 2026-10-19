"""
Appendix for the pilot paper (APA 7 style, Times New Roman).

  Appendix A  VI-SPDAT 2.0 items, scoring rules (quoted from the published tool), and how
              many of the 100 profiles meet each
  Appendix B  the exact prompt sent to Gemini for every profile
  Appendix C  the 100 profiles: demographics, VI-SPDAT score and priority, items met

Run:  python pilot/scripts/export_appendix.py
Writes: pilot/exports/pilot_appendix.docx
"""

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from analyze import apa_table, para, style_doc  # noqa: E402
from export_review_sheet import CITY, ORDER, RACE, SHORT  # noqa: E402
from gemini_prompt import INSTRUCTIONS, REPLY_FORMAT, SECTION_TITLES  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
PILOT = ROOT / "pilot"


def heading(doc, text, sub):
    para(doc, text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    para(doc, sub, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)


def main():
    inst = json.loads((ROOT / "data" / "vispdat_instrument.json").read_text(encoding="utf-8"))
    profiles = json.loads((PILOT / "data" / "profiles.json").read_text(encoding="utf-8"))["profiles"]
    n = len(profiles)
    doc = Document()
    style_doc(doc)

    # Appendix A
    heading(doc, "Appendix A", "VI-SPDAT Items, Scoring Rules, and Frequency in the Synthetic Population")
    body, groups = [], []
    for d in inst["domains"]:
        groups.append(len(body))
        title = d["name"] if d["key"] == "pre" else f"{d['letter']}. {d['name']}"
        body.append([f"{title} (max {d['max']})", "", "", ""])
        for i in d["indicators"]:
            k = sum(i["id"] in p["indicators"] for p in profiles)
            body.append(["    " + i["name"], i["questions"], i["rule"], f"{k} ({100 * k / n:.0f}%)"])
    apa_table(doc, "A1", "VI-SPDAT 2.0 (Single Adults) Scoring Items",
              ["Item", "Questions", "Scoring rule", "Profiles meeting it, n (%)"], body, size=9,
              widths=[1.7, 0.8, 3.0, 1.0], group_rows=groups,
              note=f"Rules are quoted from the VI-SPDAT American Version 2.0 for single adults "
                   f"({inst['copyright']}). Each item met scores 1 point; maximum 17. N = {n}.")
    body = []
    pri = {1: "Low", 2: "Medium", 3: "High"}
    for b in inst["bands"]:
        k = sum(b["min"] <= len(p["indicators"]) <= b["max"] for p in profiles)
        body.append([f"{b['min']}-{b['max']}", b["action"], pri[b["id"]], f"{k} ({100 * k / n:.0f}%)"])
    apa_table(doc, "A2", "VI-SPDAT Score Ranges, Recommendations, and Priority Categories Used for Comparison",
              ["Score", "VI-SPDAT recommendation", "Priority category", "Profiles, n (%)"], body, size=10,
              widths=[0.8, 3.3, 1.2, 1.2],
              note="Gemini chose among the same three priority categories, defined by the kind of help needed "
                   "(Appendix B), so the two approaches are compared category to category.")

    # Appendix B
    doc.add_page_break()
    heading(doc, "Appendix B", "Prompt Sent to Gemini")
    para(doc, "The same text was sent for every profile, one profile per request, with no other context. Only "
              "the case summary differed. Gemini was not shown VI-SPDAT questions, items, rules, or scores.",
         size=11, space_after=8)
    sample = next(p for p in profiles if p["id"] == "P01")
    blocks = [INSTRUCTIONS, "Case summary\n[Age, gender, race, and location]\n\n" + "\n".join(
        f"{t}: [text]" for _, t in SECTION_TITLES), REPLY_FORMAT]
    for block in blocks:
        for line in block.split("\n"):
            q = doc.add_paragraph()
            q.paragraph_format.space_after = Pt(0)
            q.paragraph_format.left_indent = Inches(0.5)
            r = q.add_run(line)
            r.font.name, r.font.size = "Courier New", Pt(9.5)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    para(doc, f"Example demographic line: {sample['demographic_line']}", italic=True, size=11)

    # Appendix C (landscape)
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    heading(doc, "Appendix C", "The 100 Synthetic Profiles")
    body = []
    for p in sorted(profiles, key=lambda p: (len(p["indicators"]), int(p["id"][1:]))):
        s = len(p["indicators"])
        body.append([p["id"], p["age"], "F" if p["gender"] == "Female" else "M", RACE[p["race"]],
                     CITY[p["location"]], s, "High" if s >= 8 else "Medium" if s >= 4 else "Low",
                     ", ".join(SHORT[i] for i in ORDER if i in p["indicators"]) or "none"])
    apa_table(doc, "C1", "Profiles by VI-SPDAT Score",
              ["ID", "Age", "Gender", "Race", "City", "Score", "Priority", "VI-SPDAT items met"], body, size=8,
              widths=[0.45, 0.4, 0.55, 0.6, 0.45, 0.45, 0.6, 5.4],
              note="F = female; M = male. AIAN = American Indian, Alaska Native, or Indigenous; Asian = Asian or "
                   "Asian American; Black = Black, African American, or African; Hisp = Hispanic/Latino/e/a. "
                   "Cities: LA = Los Angeles, Cin = Cincinnati, Chi = Chicago, Atl = Atlanta, NYC = New York City. "
                   "Item names are shortened from Table A1. Full case summaries are available on request.")
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for q in c.paragraphs:
                    q.paragraph_format.space_after = Pt(0)
    out = PILOT / "exports" / "pilot_appendix.docx"
    doc.save(out)
    print(f"wrote {out.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
