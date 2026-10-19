"""
Pilot analysis: composition of the 100 profiles, and VI-SPDAT vs Gemini priority.

Both approaches put each person in one of three priority categories (High / Medium /
Low; see score_vispdat.py and gemini_prompt.py), so they are compared category to
category, never by subtracting raw scores. VI-SPDAT is the comparison benchmark,
not ground truth. Every table reports counts and within-group percentages.

Run while Gemini is still in progress and it reports on the profiles done so far,
marked as preliminary.

Run:  python pilot/scripts/analyze.py
Writes: pilot/results/*.csv, pilot/exports/pilot_results.docx,
        pilot/exports/profiles_for_review.docx, pilot/results/profiles_for_review.csv
"""

import csv
import json
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
PILOT = ROOT / "pilot"
RESULTS = PILOT / "results"
EXPORTS = PILOT / "exports"
CATS = ["High", "Medium", "Low"]
RANK = {"Low": 0, "Medium": 1, "High": 2}

ITEMS = [
    ("pre_age", "Age 60 or older"),
    ("a1_sleep", "Sleeping somewhere other than a shelter or transitional housing"),
    ("a2_duration", "Homeless 1+ year continuously or 4+ episodes in 3 years"),
    ("b1_emergency", "4+ emergency service contacts in 6 months"),
    ("b2_harm", "Attacked since becoming homeless, or threatened harm to self or others"),
    ("b3_legal", "Legal issues that could mean jail or fines"),
    ("b4_exploitation", "Forced, tricked, or trading sex/drugs/favors"),
    ("c1_money", "No regular income, or owes money"),
    ("c2_activity", "No meaningful daily activity"),
    ("c3_selfcare", "Cannot meet basic needs (washing, meals)"),
    ("c4_relationships", "Homeless because of a relationship or family breakdown"),
    ("d1_physical", "Physical health condition (chronic illness, disability, pregnancy, HIV)"),
    ("d2_substance", "Substance use that cost or threatens housing"),
    ("d3_mental", "Mental health condition, head injury, or cognitive impairment"),
    ("d4_trimorbidity", "Tri-morbidity (physical + substance + mental)"),
    ("d5_medications", "Not taking, or selling, prescribed medication"),
    ("d6_trauma", "Homeless because of abuse or trauma"),
]

RACE_SHORT = {
    "American Indian, Alaska Native, or Indigenous": "American Indian, Alaska Native, or Indigenous",
    "Asian or Asian American": "Asian or Asian American",
    "Black, African American, or African": "Black, African American, or African",
    "Hispanic/Latino/e/a": "Hispanic/Latino/e/a",
    "White": "White",
}


def age_group(age):
    if age < 25:
        return "18-24"
    if age < 40:
        return "25-39"
    if age < 60:
        return "40-59"
    return "60 and older"


AGE_ORDER = ["18-24", "25-39", "40-59", "60 and older"]


def load():
    profiles = {p["id"]: p for p in json.loads((PILOT / "data" / "profiles.json").read_text(encoding="utf-8"))["profiles"]}
    with (PILOT / "data" / "vispdat_scores.csv").open(encoding="utf-8") as fh:
        vis = {r["id"]: r for r in csv.DictReader(fh)}
    gem = {}
    for f in sorted((PILOT / "data" / "gemini" / "raw").glob("P*.json")):
        rec = json.loads(f.read_text(encoding="utf-8"))
        if rec.get("priority"):
            gem[rec["id"]] = rec
    rows = []
    for pid, p in profiles.items():
        v = vis[pid]
        flags = set(p["indicators"])
        rows.append({
            "id": pid, "label": p["label"], "age": p["age"], "age_group": age_group(p["age"]),
            "gender": p["gender"], "race": p["race"], "location": p["location"],
            "flags": flags, "vispdat_score": int(v["score"]), "vispdat": v["priority"],
            "gemini": gem[pid]["priority"] if pid in gem else None,
            "explanation": gem[pid]["explanation"] if pid in gem else None,
            "model": gem[pid]["model_version"] if pid in gem else None,
            "sleeping": "Unsheltered or other (car, outdoors, couch)" if "a1_sleep" in flags else "Shelter or transitional housing",
            "mental": "Yes" if "d3_mental" in flags else "No",
            "substance": "Yes" if "d2_substance" in flags else "No",
            "physical": "Yes" if "d1_physical" in flags else "No",
            "chronic": "Yes" if "a2_duration" in flags else "No",
        })
    return profiles, rows


def pct(k, n):
    return f"{100 * k / n:.1f}" if n else "-"


# ---------------------------------------------------------------- composition

def composition(rows):
    n = len(rows)
    out = []
    for title, key, order in [
        ("Gender", "gender", ["Female", "Male"]),
        ("Race", "race", list(RACE_SHORT)),
        ("Location", "location", ["Los Angeles, California", "Cincinnati, Ohio", "Chicago, Illinois",
                                  "Atlanta, Georgia", "New York City, New York"]),
        ("Age group", "age_group", AGE_ORDER),
        ("VI-SPDAT priority", "vispdat", CATS),
    ]:
        c = Counter(r[key] for r in rows)
        for g in order:
            out.append({"characteristic": title, "group": g, "n": c[g], "percent": pct(c[g], n)})
    for ind, name in ITEMS:
        k = sum(ind in r["flags"] for r in rows)
        out.append({"characteristic": "VI-SPDAT circumstance", "group": name, "n": k, "percent": pct(k, n)})
    return out


# ---------------------------------------------------------------- comparison

COMPARE = [
    ("Gender", "gender", ["Female", "Male"]),
    ("Race", "race", list(RACE_SHORT)),
    ("Location", "location", ["Los Angeles, California", "Cincinnati, Ohio", "Chicago, Illinois",
                              "Atlanta, Georgia", "New York City, New York"]),
    ("Age group", "age_group", AGE_ORDER),
    ("Current sleeping situation", "sleeping", ["Shelter or transitional housing",
                                                "Unsheltered or other (car, outdoors, couch)"]),
    ("Homeless 1+ year or 4+ episodes", "chronic", ["Yes", "No"]),
    ("Physical health condition", "physical", ["Yes", "No"]),
    ("Mental health condition", "mental", ["Yes", "No"]),
    ("Substance use", "substance", ["Yes", "No"]),
]


def group_row(title, g, rs):
    n = len(rs)
    row = {"characteristic": title, "group": g, "n": n}
    for tool in ("vispdat", "gemini"):
        c = Counter(r[tool] for r in rs)
        for cat in CATS:
            row[f"{tool}_{cat.lower()}_n"] = c[cat]
            row[f"{tool}_{cat.lower()}_pct"] = pct(c[cat], n)
    same = sum(r["vispdat"] == r["gemini"] for r in rs)
    higher = sum(RANK[r["gemini"]] > RANK[r["vispdat"]] for r in rs)
    lower = sum(RANK[r["gemini"]] < RANK[r["vispdat"]] for r in rs)
    row.update({"same_n": same, "same_pct": pct(same, n), "gemini_higher_n": higher,
                "gemini_higher_pct": pct(higher, n), "gemini_lower_n": lower, "gemini_lower_pct": pct(lower, n)})
    return row


def comparisons(done):
    out = [group_row("All profiles", "All", done)]
    for title, key, order in COMPARE:
        for g in order:
            out.append(group_row(title, g, [r for r in done if r[key] == g]))
    return out


def crosstab(done):
    return {(v, g): sum(r["vispdat"] == v and r["gemini"] == g for r in done) for v in CATS for g in CATS}


def weighted_kappa(done):
    """Linear-weighted Cohen's kappa between the two ordered category sets."""
    n = len(done)
    if not n:
        return None
    k = len(CATS)
    obs = crosstab(done)
    rv = Counter(r["vispdat"] for r in done)
    rg = Counter(r["gemini"] for r in done)
    num = den = 0.0
    for v in CATS:
        for g in CATS:
            w = abs(RANK[v] - RANK[g]) / (k - 1)
            num += w * obs[(v, g)] / n
            den += w * (rv[v] / n) * (rg[g] / n)
    return 1 - num / den if den else None


def write_csv(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)


# ---------------------------------------------------------------- Word export (APA 7 style)

def style_doc(doc, landscape=False):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    sec = doc.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(1))


def border(cell, edges):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for e in ("top", "bottom", "left", "right"):
        el = OxmlElement(f"w:{e}")
        if e in edges:
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "6"); el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
        b.append(el)
    tcPr.append(b)


def para(doc, text, bold=False, italic=False, size=12, align=None, space_after=0):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic, r.font.size = bold, italic, Pt(size)
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    return p


def apa_table(doc, number, title, header, body, note=None, size=10, widths=None, group_rows=()):
    para(doc, f"Table {number}", bold=True)
    para(doc, title, italic=True, space_after=4)
    t = doc.add_table(rows=1 + len(body), cols=len(header))
    t.autofit = widths is None
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.font.size = Pt(size)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
        border(c, {"top", "bottom"})
    for i, row in enumerate(body, start=1):
        for j, v in enumerate(row):
            c = t.rows[i].cells[j]
            c.text = ""
            run = c.paragraphs[0].add_run(str(v))
            run.font.size = Pt(size)
            if i - 1 in group_rows:
                run.italic = True
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            border(c, {"bottom"} if i == len(body) else set())
    if widths:
        for row in t.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Inches(w)
    if note:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        r = p.add_run("Note. ")
        r.italic, r.font.size = True, Pt(size)
        r2 = p.add_run(note)
        r2.font.size = Pt(size)
    doc.add_paragraph()


def export_results(rows, done, comp, comps, ct, kappa):
    doc = Document()
    style_doc(doc, landscape=True)
    n, nd = len(rows), len(done)
    status = "" if nd == n else f" (PRELIMINARY: Gemini has assessed {nd} of {n} profiles)"
    model = sorted({r["model"] for r in done})
    para(doc, "Pilot Results: VI-SPDAT and Gemini Housing Priority" + status, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    para(doc, f"Gemini model: {', '.join(model)}. One independent request per person, identical instructions; "
              "Gemini was not shown VI-SPDAT questions or scores.", size=11, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_after=12)

    # Table 1 composition
    body, groups = [], []
    last = None
    for r in comp:
        if r["characteristic"] != last:
            groups.append(len(body))
            body.append([r["characteristic"], "", ""])
            last = r["characteristic"]
        body.append(["    " + r["group"], r["n"], r["percent"]])
    apa_table(doc, 1, "Composition of the Synthetic Population (N = 100)", ["Characteristic", "n", "%"], body,
              note="VI-SPDAT circumstances are the 17 items of VI-SPDAT 2.0 for single adults; one person can "
                   "have several. Gender, race, and location were assigned so that each race has 10 women, 10 men, "
                   "and 4 people in each city, with a similar mix of VI-SPDAT need.",
              widths=[5.2, 0.8, 0.8], group_rows=groups)

    # Table 2 cross-tab
    body = []
    for v in CATS:
        tot = sum(ct[(v, g)] for g in CATS)
        body.append([f"VI-SPDAT {v}"] + [ct[(v, g)] for g in CATS] + [tot])
    body.append(["Total"] + [sum(ct[(v, g)] for v in CATS) for g in CATS] + [nd])
    same = sum(ct[(c, c)] for c in CATS)
    higher = sum(ct[(v, g)] for v in CATS for g in CATS if RANK[g] > RANK[v])
    lower = sum(ct[(v, g)] for v in CATS for g in CATS if RANK[g] < RANK[v])
    apa_table(doc, 2, "Priority Assigned by VI-SPDAT and by Gemini to the Same Profiles",
              ["", "Gemini High", "Gemini Medium", "Gemini Low", "Total"], body,
              note=f"Same category: {same} of {nd} ({pct(same, nd)}%). Gemini higher: {higher} ({pct(higher, nd)}%). "
                   f"Gemini lower: {lower} ({pct(lower, nd)}%). "
                   + (f"Linear-weighted Cohen's kappa = {kappa:.2f}. " if kappa is not None else "") +
                   "VI-SPDAT categories: High = score 8-17 (Permanent Supportive Housing), Medium = 4-7 "
                   "(Rapid Re-Housing), Low = 0-3 (no housing intervention).",
              widths=[1.8, 1.2, 1.2, 1.2, 1.0])

    # Table 3 by group
    header = ["Group", "n", "VI-SPDAT High\nn (%)", "Gemini High\nn (%)", "VI-SPDAT Medium\nn (%)",
              "Gemini Medium\nn (%)", "VI-SPDAT Low\nn (%)", "Gemini Low\nn (%)", "Same\nn (%)",
              "Gemini higher\nn (%)", "Gemini lower\nn (%)"]
    body, groups, last = [], [], None
    for r in comps:
        if r["characteristic"] != last:
            if r["characteristic"] != "All profiles":
                groups.append(len(body))
                body.append([r["characteristic"]] + [""] * 10)
            last = r["characteristic"]
        cell = lambda k: f"{r[k + '_n']} ({r[k + '_pct']})"
        label = r["group"] if r["characteristic"] == "All profiles" else "    " + r["group"]
        body.append([label, r["n"], cell("vispdat_high"), cell("gemini_high"), cell("vispdat_medium"),
                     cell("gemini_medium"), cell("vispdat_low"), cell("gemini_low"), cell("same"),
                     cell("gemini_higher"), cell("gemini_lower")])
    apa_table(doc, 3, "Priority by Group: VI-SPDAT and Gemini", header, body, size=8,
              note="Percentages are within each group (row). Same, higher, and lower compare the two "
                   "categories for each person; higher means Gemini placed the person in a higher priority "
                   "category than VI-SPDAT. Differences between groups describe the two approaches and do not "
                   "by themselves establish bias, since groups can differ in circumstances.",
              widths=[2.3, 0.4] + [0.65] * 9, group_rows=groups)

    out = EXPORTS / "pilot_results.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


def export_review(profiles, rows):
    doc = Document()
    style_doc(doc)
    para(doc, "Pilot Profiles for Review (N = 100)", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    para(doc, "Each profile shows the demographic line and the case summary exactly as Gemini receives them, "
              "followed by the VI-SPDAT items it meets and its VI-SPDAT score. Check that every person differs, "
              "that the story is plausible and internally consistent, and that the items match the text.",
         size=11, space_after=12)
    names = dict(ITEMS)
    from gemini_prompt import SECTION_TITLES
    csv_rows = []
    for r in rows:
        p = profiles[r["id"]]
        para(doc, f"{r['id']}  {p['label']}", bold=True, space_after=2)
        para(doc, p["demographic_line"], italic=True, size=11, space_after=4)
        for key, title in SECTION_TITLES:
            q = doc.add_paragraph()
            q.paragraph_format.space_after = Pt(2)
            a = q.add_run(f"{title}: ")
            a.bold, a.font.size = True, Pt(11)
            b = q.add_run(p["narrative"][key])
            b.font.size = Pt(11)
        items = [names[i] for i, _ in ITEMS if i in r["flags"]]
        para(doc, f"VI-SPDAT: {r['vispdat_score']} / 17, {r['vispdat']} priority. Items met: "
                  + ("; ".join(items) if items else "none"), size=10, space_after=12)
        csv_rows.append({"id": r["id"], "label": p["label"], "source": p["source"], "age": p["age"],
                         "gender": p["gender"], "race": p["race"], "location": p["location"],
                         "vispdat_score": r["vispdat_score"], "vispdat_priority": r["vispdat"],
                         "items_met": "; ".join(items),
                         **{title: p["narrative"][key] for key, title in SECTION_TITLES}})
    out = EXPORTS / "profiles_for_review.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    write_csv(RESULTS / "profiles_for_review.csv", csv_rows)
    return out


def main():
    profiles, rows = load()
    done = [r for r in rows if r["gemini"]]
    comp = composition(rows)
    write_csv(RESULTS / "composition.csv", comp)
    write_csv(RESULTS / "by_profile.csv", [
        {k: r[k] for k in ("id", "age", "gender", "race", "location", "vispdat_score", "vispdat", "gemini",
                           "explanation")} for r in rows])
    try:
        review = export_review(profiles, rows)
        print(f"wrote {review.relative_to(ROOT)} and pilot/results/composition.csv")
    except PermissionError:
        print("profiles_for_review.docx is open in Word; close it and rerun to refresh it")
    if not done:
        print("no Gemini replies yet; comparison tables are written once replies exist")
        return
    comps = comparisons(done)
    write_csv(RESULTS / "comparison_by_group.csv", comps)
    ct = crosstab(done)
    kappa = weighted_kappa(done)
    out = export_results(rows, done, comp, comps, ct, kappa)
    print(f"wrote {out.relative_to(ROOT)} ({len(done)}/{len(rows)} profiles assessed by Gemini)")
    same = sum(ct[(c, c)] for c in CATS)
    print(f"  same category {same}/{len(done)}; weighted kappa {kappa:.2f}" if kappa is not None else "")


if __name__ == "__main__":
    import sys
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    main()
